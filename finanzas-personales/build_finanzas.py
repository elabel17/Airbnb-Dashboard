# -*- coding: utf-8 -*-
"""
Generador del sistema de Finanzas Personales (Excel 365, sin macros).

Uso:  python build_finanzas.py [ruta_salida.xlsx]

Todo el libro se construye con fórmulas nativas (SUMIFS, INDEX/MATCH, LARGE,
COUNTIFS, EOMONTH...) para que cada métrica se recalcule al agregar movimientos.
"""
import random
import sys
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.chart import BarChart, DoughnutChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.chart.series import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText
from openpyxl.chart.axis import ChartLines
from openpyxl.comments import Comment
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import CharacterProperties, Font as DFont, Paragraph, ParagraphProperties
from openpyxl.formatting.rule import CellIsRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import column_index_from_string, get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.workbook.defined_name import DefinedName

OUT = sys.argv[1] if len(sys.argv) > 1 else "Finanzas_Personales.xlsx"

# ----------------------------------------------------------------------------
# Diseño
# ----------------------------------------------------------------------------
FONT = "Aptos"
INK, TEXT, MUTED, LINE_C = "0F172A", "334155", "64748B", "E2E8F0"
PAGE, CARD, SOFT = "F1F5F9", "FFFFFF", "F8FAFC"
BLUE, ORANGE, AQUA, YELLOW, VIOLET = "2A78D6", "EB6834", "1BAF7A", "EDA100", "6250D6"
GOOD, GOOD_BG = "15803D", "DCFCE7"
WARN, WARN_BG = "B45309", "FEF3C7"
BAD, BAD_BG = "B91C1C", "FEE2E2"
INPUT_BG, INPUT_LINE = "FFF8E7", "F2C46D"
DEMO_BG = "F3F0FF"

MONEY = '"RD$ "#,##0;[Red]-"RD$ "#,##0;"–"'
MONEY2 = '#,##0.00;[Red]-#,##0.00;"–"'
NUM = '#,##0;[Red]-#,##0;"–"'
PCT = '0.0%;[Red]-0.0%;"–"'
DATE = "dd/mm/yyyy"

thin = Side(style="thin", color=LINE_C)
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
BOTTOM = Border(bottom=thin)
inp_side = Side(style="thin", color=INPUT_LINE)
INPUT_BOX = Border(left=inp_side, right=inp_side, top=inp_side, bottom=inp_side)


def F(size=10, bold=False, color=TEXT, italic=False, underline=None):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic, underline=underline)


def fill(c):
    return PatternFill("solid", start_color=c, end_color=c)


def put(ws, ref, value=None, font=None, fill_c=None, fmt=None, align=None, border=None):
    c = ws[ref]
    if value is not None:
        c.value = value
    if font is not None:
        c.font = font
    if fill_c is not None:
        c.fill = fill(fill_c)
    if fmt is not None:
        c.number_format = fmt
    if align is not None:
        c.alignment = align
    if border is not None:
        c.border = border
    return c


def area(ws, rng, fill_c=None, border=None, font=None, align=None, fmt=None):
    for row in ws[rng]:
        for c in row:
            if fill_c is not None:
                c.fill = fill(fill_c)
            if border is not None:
                c.border = border
            if font is not None:
                c.font = font
            if align is not None:
                c.alignment = align
            if fmt is not None:
                c.number_format = fmt


def merge(ws, rng, value=None, **kw):
    ws.merge_cells(rng)
    first = rng.split(":")[0]
    return put(ws, first, value, **kw)


def outline_box(ws, rng, color=LINE_C):
    """Borde exterior de un rango (tarjetas)."""
    s = Side(style="thin", color=color)
    rows = list(ws[rng])
    for i, row in enumerate(rows):
        for j, c in enumerate(row):
            c.border = Border(
                left=s if j == 0 else c.border.left,
                right=s if j == len(row) - 1 else c.border.right,
                top=s if i == 0 else c.border.top,
                bottom=s if i == len(rows) - 1 else c.border.bottom,
            )


LEFT = Alignment(horizontal="left", vertical="center", indent=1)
LEFT_WRAP = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
TOP_WRAP = Alignment(horizontal="left", vertical="top", wrap_text=True, indent=1)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="center", indent=1)

SHEETS = [
    ("Inicio", "Inicio", INK),
    ("Registro", "Registro", BLUE),
    ("Dashboard", "Dashboard", BLUE),
    ("Alimentación", "Alimentación", ORANGE),
    ("Gastos hormiga", "Hormiga", ORANGE),
    ("Recurrentes", "Recurrentes", ORANGE),
    ("Presupuestos", "Presupuestos", AQUA),
    ("Cuentas", "Cuentas", AQUA),
    ("Categorías", "Categorías", MUTED),
    ("Config", "Config", MUTED),
    ("Motor", "Motor", "94A3B8"),
]


def q(sheet):
    return "'" + sheet + "'"


def link(ws, ref, sheet, text, cell="A1", font=None):
    c = ws[ref]
    c.value = text
    c.hyperlink = Hyperlink(ref=ref, location=f"{q(sheet)}!{cell}", display=text)
    c.font = font or F(9, True, BLUE)
    return c


def header(ws, title, subtitle, last_col, nav_row=2, nav_start=2, nav_span=1):
    """Banda superior + barra de navegación."""
    ws.sheet_view.showGridLines = False
    ws.row_dimensions[1].height = 38
    ws.row_dimensions[2].height = 22
    area(ws, f"A1:{L(last_col)}1", fill_c=INK)
    put(ws, "B1", title, F(18, True, "FFFFFF"), align=Alignment(vertical="center"))
    sub_col = max(6, last_col - 6)
    put(ws, f"{L(sub_col)}1", subtitle, F(9, False, "CBD5E1"),
        align=Alignment(horizontal="right", vertical="center"))
    if sub_col != last_col:
        ws.merge_cells(f"{L(sub_col)}1:{L(last_col)}1")
    area(ws, f"A2:{L(last_col)}2", fill_c="E8EEF6")
    NAV_LAST[ws.title] = last_col


NAV_LAST = {}


def nav(ws):
    """Barra de navegación: coloca cada enlace donde quepa según el ancho real de las columnas."""
    def width(ci):
        d = ws.column_dimensions.get(L(ci))
        if d is None:
            return 8.43
        if d.hidden:
            return 0
        return d.width if d.customWidth else 8.43
    col = 2
    for sheet, short, _ in SHEETS:
        if sheet == "Motor":
            continue
        while width(col) < 2:
            col += 1
        is_me = sheet == ws.title
        text = ("● " if is_me else "") + short
        c = link(ws, f"{L(col)}2", sheet, text,
                 font=F(9, True, INK if is_me else BLUE, underline=None if is_me else "single"))
        c.alignment = Alignment(horizontal="left", vertical="center")
        need, acc = len(text) * 0.95 + 1.5, 0
        while acc < need:
            acc += width(col)
            col += 1
    for ci in range(2, max(col, NAV_LAST.get(ws.title, col)) + 1):
        ws.cell(row=2, column=ci).fill = fill("E8EEF6")


def section(ws, ref_range, text, color=INK, size=12):
    first = ref_range.split(":")[0]
    if ":" in ref_range:
        ws.merge_cells(ref_range)
    put(ws, first, text, F(size, True, color), align=Alignment(vertical="bottom"))
    for row in ws[ref_range]:
        for c in row:
            c.border = Border(bottom=Side(style="medium", color=color))


def note(ws, ref, text, size=9, color=MUTED, italic=True, wrap=False):
    return put(ws, ref, text, F(size, False, color, italic=italic),
               align=Alignment(vertical="center", wrap_text=wrap))


def input_cell(ws, ref, value=None, fmt=None, align=None):
    c = put(ws, ref, value, F(10, True, "1E3A8A"), INPUT_BG, fmt, align or CENTER, INPUT_BOX)
    return c


def add_name(wb, name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)


def axis_text(size=900, color="475569"):
    cp = CharacterProperties(latin=DFont(typeface=FONT), sz=size, solidFill=color)
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])


def style_chart(ch, legend=True, gridlines=True):
    ch.style = 2
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.x_axis.txPr = axis_text()
    ch.y_axis.txPr = axis_text()
    ch.x_axis.majorTickMark = "none"
    ch.y_axis.majorTickMark = "none"
    ch.x_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill=LINE_C))
    ch.y_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))
    if gridlines:
        ch.y_axis.majorGridlines = ChartLines(spPr=GraphicalProperties(ln=LineProperties(solidFill="EEF2F7")))
    else:
        ch.y_axis.majorGridlines = None
    if legend:
        ch.legend.position = "t"
        ch.legend.txPr = axis_text(900, TEXT)
    else:
        ch.legend = None
    try:
        ch.graphical_properties = GraphicalProperties(ln=LineProperties(noFill=True))
    except Exception:
        pass


def color_series(s, color, line=False, width=28575):
    if line:
        s.graphicalProperties.line.solidFill = color
        s.graphicalProperties.line.width = width
        s.smooth = False
        s.marker.symbol = "circle"
        s.marker.size = 6
        s.marker.graphicalProperties = GraphicalProperties(solidFill=color)
        s.marker.graphicalProperties.line.solidFill = "FFFFFF"
    else:
        s.graphicalProperties.solidFill = color
        s.graphicalProperties.line.solidFill = "FFFFFF"


def data_labels(s, fmt="#,##0", pos=None):
    s.dLbls = DataLabelList()
    s.dLbls.showVal = True
    s.dLbls.showSerName = False
    s.dLbls.showCatName = False
    s.dLbls.showLegendKey = False
    s.dLbls.numFmt = fmt
    s.dLbls.txPr = axis_text(800, TEXT)
    if pos:
        s.dLbls.position = pos


MESES_ES = '"Ene","Feb","Mar","Abr","May","Jun","Jul","Ago","Sep","Oct","Nov","Dic"'


def fecha_txt(x):
    """Fecha como texto dd/mm/aaaa sin TEXT() (evita códigos de formato dependientes del idioma)."""
    return f'RIGHT("0"&DAY({x}),2)&"/"&RIGHT("0"&MONTH({x}),2)&"/"&YEAR({x})'


# ----------------------------------------------------------------------------
# Constantes de estructura
# ----------------------------------------------------------------------------
R0, NROWS = 6, 5000
RL = R0 + NROWS - 1           # última fila de la tabla de movimientos
CAT0, CATN = 6, 200           # filas de categorías
CATL = CAT0 + CATN - 1
CTA0, CTAN = 7, 30            # filas de cuentas
CTAL = CTA0 + CTAN - 1
REC0, RECN = 7, 50            # filas de recurrentes
RECL = REC0 + RECN - 1

wb = Workbook()
wb.remove(wb.active)
ws_ = {}
for name, _, color in SHEETS:
    w = wb.create_sheet(name)
    w.sheet_properties.tabColor = color
    ws_[name] = w

INI, REG, DSH, ALI, HOR, REC, PRE, CTA, CAT, CFG, MOT = (ws_[s[0]] for s in SHEETS)

# ----------------------------------------------------------------------------
# Nombres definidos de columnas del Registro
# ----------------------------------------------------------------------------
REG_COLS = {
    "rFecha": "B", "rTipo": "C", "rCat": "D", "rSub": "E", "rDesc": "F", "rMonto": "G",
    "rMon": "H", "rTasa": "I", "rCta": "J", "rDest": "K", "rCom": "L", "rClas": "M",
    "rObs": "N", "rEstado": "O", "rTipoE": "P", "rCtaE": "Q", "rTasaRef": "R",
    "rDOP": "S", "rUSD": "T", "rGasto": "U", "rIng": "V", "rTrad": "W", "rClasE": "X",
    "rGrupo": "Y", "rDia": "Z", "rHorm": "AA", "rCtaSO": "AB", "rCtaSD": "AC",
    "rTCO": "AD", "rTCD": "AE", "rEfO": "AF", "rEfD": "AG", "rFlujo": "AH", "rComIdx": "AI",
}
for nm, col in REG_COLS.items():
    add_name(wb, nm, f"'Registro'!${col}${R0}:${col}${RL}")

CAT_COLS = {"catCat": "B", "catSub": "C", "catTipo": "D", "catClasif": "E", "catGrupo": "F",
            "catComida": "G", "catClave": "P", "catIdx": "Q", "catFoodIdx": "R"}
for nm, col in CAT_COLS.items():
    add_name(wb, nm, f"'Categorías'!${col}${CAT0}:${col}${CATL}")
add_name(wb, "catUList", f"'Categorías'!$L${CAT0}:$L${CAT0 + 39}")
add_name(wb, "lstCat", f"OFFSET('Categorías'!$L${CAT0},0,0,MAX(1,COUNTIF('Categorías'!$L${CAT0}:$L${CAT0 + 39},\"?*\")),1)")

CTA_COLS = {"ctaNombre": "B", "ctaTipo": "C", "ctaMoneda": "D", "ctaVinc": "E", "ctaIni": "F",
            "ctaFechaIni": "G", "ctaLiq": "I", "ctaMov": "J", "ctaSaldo": "K", "ctaSaldoDOP": "L"}
for nm, col in CTA_COLS.items():
    add_name(wb, nm, f"'Cuentas'!${col}${CTA0}:${col}${CTAL}")
add_name(wb, "lstCuentas", f"OFFSET('Cuentas'!$B${CTA0},0,0,MAX(1,COUNTIF('Cuentas'!$B${CTA0}:$B${CTAL},\"?*\")),1)")

# ============================================================================
# CONFIG
# ============================================================================
ws = CFG
header(ws, "Configuración", "Parámetros y listas que alimentan los desplegables", 33)
for col, w in {"A": 2, "B": 30, "C": 16, "D": 13, "E": 13, "F": 62, "G": 3}.items():
    ws.column_dimensions[col].width = w
section(ws, "B4:F4", "Parámetros generales")
params = [
    (5, "Tasa USD → DOP (referencial)", 63.50, "0.00",
     "Valor referencial de ejemplo: ACTUALÍZALO con la tasa vigente. Se usa cuando un movimiento en USD no trae su propia tasa."),
    (6, "Cuenta predeterminada", None, None,
     "Opcional. Si la llenas, los movimientos sin cuenta usarán esta. Déjala vacía para que el sistema te avise cuando falte."),
    (7, "Categoría de alimentación", "Alimentación", None,
     "Nombre exacto de la categoría que analiza la hoja Alimentación."),
    (8, "Categoría de trading", "Trading y actividades financieras", None,
     "Nombre exacto de la categoría cuyos gastos se reportan como costos operativos de trading."),
    (9, "Umbral gasto hormiga", "=Umbral", None, "Se configura en la hoja Gastos hormiga."),
]
for r, label, val, fmt, nt in params:
    put(ws, f"B{r}", label, F(10, True, INK), align=LEFT)
    if r == 9:
        put(ws, f"C{r}", val, F(10, False, MUTED), fmt=MONEY, align=CENTER)
    else:
        input_cell(ws, f"C{r}", val, fmt)
    ws.merge_cells(f"D{r}:F{r}")
    note(ws, f"D{r}", nt, wrap=True)
    ws.row_dimensions[r].height = 30
add_name(wb, "TasaUSD", "'Config'!$C$5")
add_name(wb, "CuentaPred", "'Config'!$C$6")
add_name(wb, "catAlim", "'Config'!$C$7")
add_name(wb, "catTrading", "'Config'!$C$8")
ws["C5"].comment = Comment("Tasa referencial de ejemplo, no oficial. Reemplázala por la tasa vigente.", "Sistema")

section(ws, "B11:F11", "Tipos de movimiento · cómo se contabiliza cada uno")
tipos = [
    ("Gasto", "Sí", "No", "Sale de la cuenta", "Consumo real. Una compra con tarjeta de crédito se registra AQUÍ, en la fecha de consumo, con la tarjeta como cuenta."),
    ("Ingreso", "No", "Sí", "Entra a la cuenta", "Salario, ingresos extra, intereses. Un retiro desde trading NO es ingreso: es Transferencia."),
    ("Transferencia", "No", "No", "Origen → destino", "Movimiento entre cuentas propias (banco→efectivo, banco→trading, trading→banco). No es gasto ni ingreso."),
    ("Pago tarjeta", "No", "No", "Origen → tarjeta", "Pago/liquidación de la tarjeta de crédito. Reduce la deuda; NO es un segundo gasto."),
    ("Reembolso", "Resta", "No", "Entra a la cuenta", "Te devuelven dinero de un gasto. Usa la misma categoría del gasto original: lo neutraliza."),
    ("Ajuste", "No", "No", "± a la cuenta", "Corrección de saldo al conciliar (monto positivo o negativo). No afecta reportes de gasto."),
    ("Resultado trading", "No (trading)", "No (trading)", "± a la cuenta trading", "Ganancia (+) o pérdida (−) REALIZADA en una cuenta de trading. Se reporta aparte del gasto real."),
]
for j, h in enumerate(["Tipo", "¿Gasto real?", "¿Ingreso?", "Efecto en saldos", "Cómo usarlo"]):
    put(ws, f"{L(2 + j)}12", h, F(9, True, "FFFFFF"), "334155", align=CENTER if j else LEFT)
for i, t in enumerate(tipos):
    r = 13 + i
    for j, v in enumerate(t):
        put(ws, f"{L(2 + j)}{r}", v, F(10, j == 0, INK if j == 0 else TEXT),
            CARD, align=LEFT_WRAP if j in (0, 4) else CENTER, border=BOX)
    ws.row_dimensions[r].height = 32
add_name(wb, "lstTipos", "'Config'!$B$13:$B$19")

# Listas auxiliares (a la derecha)
lists = {
    "H": ("Monedas", ["DOP", "USD"]),
    "J": ("Tipos de cuenta", ["Efectivo", "Cuenta bancaria", "Tarjeta de débito", "Tarjeta de crédito", "Billetera digital", "Cuenta de trading"]),
    "K": ("¿Es liquidez?", ["Sí", "Sí", "Sí", "No", "Sí", "No"]),
    "M": ("Clasificación", ["Fijo", "Variable"]),
    "O": ("Frecuencia", ["Semanal", "Quincenal", "Mensual", "Bimestral", "Trimestral", "Semestral", "Anual"]),
    "P": ("Meses", [0, 0, 1, 2, 3, 6, 12]),
    "Q": ("Días", [7, 15, 0, 0, 0, 0, 0]),
    "R": ("× mensual", [52 / 12, 2, 1, 0.5, 1 / 3, 1 / 6, 1 / 12]),
    "T": ("Períodos", ["Mes actual", "Mes anterior", "Mes seleccionado", "Últimos 7 días", "Últimos 30 días",
                       "Últimos 90 días", "Año actual", "Año anterior", "Todo el historial", "Personalizado"]),
    "V": ("Grupos de alimentación", ["Comida preparada", "Supermercado e ingredientes", "Bebidas y snacks"]),
    "X": ("Comida del día", ["Desayuno", "Almuerzo", "Cena"]),
    "Z": ("Sí / No", ["Sí", "No"]),
    "AB": ("Tipo de categoría", ["Gasto", "Ingreso"]),
}
section(ws, "H4:AF4", "Listas de los desplegables (puedes ampliarlas con cuidado)")
for col, (title, vals) in lists.items():
    ws.column_dimensions[col].width = max(12, min(26, max(len(str(v)) for v in vals + [title]) + 3))
    put(ws, f"{col}12", title, F(9, True, "FFFFFF"), "334155", align=CENTER)
    for i, v in enumerate(vals):
        c = put(ws, f"{col}{13 + i}", v, F(10), CARD, border=BOX, align=LEFT)
        if col == "R":
            c.number_format = "0.000"
for col in ("I", "L", "N", "S", "U", "W", "Y", "AA", "AC", "AE"):
    ws.column_dimensions[col].width = 2
add_name(wb, "lstMonedas", "'Config'!$H$13:$H$14")
add_name(wb, "tcTipo", "'Config'!$J$13:$J$22")
add_name(wb, "tcLiq", "'Config'!$K$13:$K$22")
add_name(wb, "lstTiposCta", "'Config'!$J$13:$J$18")
add_name(wb, "lstClasif", "'Config'!$M$13:$M$14")
add_name(wb, "frecNombre", "'Config'!$O$13:$O$19")
add_name(wb, "frecMeses", "'Config'!$P$13:$P$19")
add_name(wb, "frecDias", "'Config'!$Q$13:$Q$19")
add_name(wb, "frecFactor", "'Config'!$R$13:$R$19")
add_name(wb, "lstPeriodos", "'Config'!$T$13:$T$22")
add_name(wb, "lstGrupos", "'Config'!$V$13:$V$15")
add_name(wb, "gCP", "'Config'!$V$13")
add_name(wb, "gSI", "'Config'!$V$14")
add_name(wb, "gBS", "'Config'!$V$15")
add_name(wb, "lstComidas", "'Config'!$X$13:$X$15")
add_name(wb, "lstSiNo", "'Config'!$Z$13:$Z$14")
add_name(wb, "lstTipoCat", "'Config'!$AB$13:$AB$14")

# Lista de meses (automática, últimos 36 meses) y filtros con "Todas"
ws.column_dimensions["AD"].width = 12
ws.column_dimensions["AF"].width = 30
ws.column_dimensions["AG"].width = 24
put(ws, "AD12", "Meses (auto)", F(9, True, "FFFFFF"), "334155", align=CENTER)
for k in range(36):
    put(ws, f"AD{13 + k}",
        f'=YEAR(DATE(YEAR(TODAY()),MONTH(TODAY())-{k},1))&"-"&RIGHT("0"&MONTH(DATE(YEAR(TODAY()),MONTH(TODAY())-{k},1)),2)',
        F(10), CARD, border=BOX, align=LEFT)
add_name(wb, "lstMeses", "'Config'!$AD$13:$AD$48")
put(ws, "AF12", "Filtro categorías (auto)", F(9, True, "FFFFFF"), "334155", align=CENTER)
put(ws, "AG12", "Filtro cuentas (auto)", F(9, True, "FFFFFF"), "334155", align=CENTER)
put(ws, "AF13", "Todas", F(10, True), CARD, border=BOX, align=LEFT)
put(ws, "AG13", "Todas", F(10, True), CARD, border=BOX, align=LEFT)
for k in range(1, 41):
    put(ws, f"AF{13 + k}", f'=IFERROR(INDEX(catUList,{k})&"","")', F(10), CARD, border=BOX, align=LEFT)
for k in range(1, 31):
    put(ws, f"AG{13 + k}", f'=IFERROR(INDEX(ctaNombre,{k})&"","")', F(10), CARD, border=BOX, align=LEFT)
add_name(wb, "lstCatFiltro", "OFFSET('Config'!$AF$13,0,0,COUNTIF('Config'!$AF$13:$AF$53,\"?*\"),1)")
add_name(wb, "lstCtaFiltro", "OFFSET('Config'!$AG$13,0,0,COUNTIF('Config'!$AG$13:$AG$43,\"?*\"),1)")
note(ws, "H24", "Si agregas un tipo de cuenta, indica en «¿Es liquidez?» si es dinero disponible (Sí) o no (No).", wrap=False)
note(ws, "H25", "No cambies los nombres de los Tipos de movimiento ni de los Períodos: las fórmulas dependen de ellos.", wrap=False)

# ============================================================================
# CATEGORÍAS
# ============================================================================
ws = CAT
header(ws, "Categorías y subcategorías", "Edita, agrega o elimina filas · se reflejan en todo el sistema", 14)
for col, w in {"A": 2, "B": 30, "C": 36, "D": 12, "E": 15, "F": 27, "G": 15, "H": 3, "I": 3, "J": 3, "K": 3, "L": 32, "M": 12, "N": 3}.items():
    ws.column_dimensions[col].width = w
for col in ("P", "Q", "R"):
    ws.column_dimensions[col].hidden = True
merge(ws, "B3:G3",
      "Mantén juntas las filas de una misma categoría (ordena por Categoría si agregas filas al final). "
      "«Clasificación sugerida» se usa cuando dejas vacía la columna Fijo/Variable del Registro. "
      "«Grupo de alimentación» alimenta el análisis especial de comida.",
      font=F(9, False, MUTED, True), align=LEFT_WRAP)
ws.row_dimensions[3].height = 30
heads = ["Categoría", "Subcategoría", "Tipo", "Clasificación sugerida", "Grupo de alimentación", "Comida del día"]
for j, h in enumerate(heads):
    put(ws, f"{L(2 + j)}5", h, F(10, True, "FFFFFF"), INK, align=LEFT)
put(ws, "L5", "Lista de categorías (automática)", F(10, True, "FFFFFF"), "475569", align=LEFT)
put(ws, "M5", "Tipo", F(10, True, "FFFFFF"), "475569", align=CENTER)

CP, SI, BS = "Comida preparada", "Supermercado e ingredientes", "Bebidas y snacks"
CATS = [
    ("Alimentación", [("Desayunos comprados", "Variable", CP, "Desayuno"),
                      ("Almuerzos laborales", "Variable", CP, "Almuerzo"),
                      ("Cenas compradas", "Variable", CP, "Cena"),
                      ("Restaurantes", "Variable", CP, None),
                      ("Delivery", "Variable", CP, None),
                      ("Supermercado", "Variable", SI, None),
                      ("Ingredientes para cocinar", "Variable", SI, None),
                      ("Bebidas, snacks y otros consumos", "Variable", BS, None)]),
    ("Hogar y servicios", [("Internet", "Fijo"), ("Electricidad", "Fijo"), ("Agua", "Fijo"),
                           ("Compras del hogar", "Variable"), ("Mantenimiento", "Variable"), ("Otros servicios", "Variable")]),
    ("Transporte y vehículo", [("Combustible", "Variable"), ("Mantenimiento", "Variable"), ("Reparaciones", "Variable"),
                               ("Seguros", "Fijo"), ("Parqueos y peajes", "Variable")]),
    ("Tecnología y suscripciones", [("Herramientas de inteligencia artificial", "Fijo"), ("Servicios digitales", "Fijo"),
                                    ("Software", "Fijo"), ("Entretenimiento", "Fijo"), ("Otras suscripciones", "Fijo")]),
    ("Trading y actividades financieras", [("Evaluaciones y cuentas de trading", "Variable"), ("Plataformas y herramientas", "Fijo"),
                                           ("Comisiones", "Variable"), ("Depósitos y retiros", "Variable"),
                                           ("Ganancias y pérdidas realizadas", "Variable")]),
    ("Salud", [("Consultas médicas", "Variable"), ("Medicamentos y farmacia", "Variable"), ("Seguro médico", "Fijo"),
               ("Gimnasio y bienestar", "Fijo")]),
    ("Educación", [("Cursos y certificaciones", "Variable"), ("Libros y materiales", "Variable"), ("Matrícula", "Fijo")]),
    ("Compras personales", [("Ropa y calzado", "Variable"), ("Cuidado personal", "Variable"), ("Electrónica", "Variable"),
                            ("Otras compras", "Variable")]),
    ("Pareja y actividades sociales", [("Salidas y citas", "Variable"), ("Actividades sociales", "Variable"),
                                       ("Viajes y paseos", "Variable")]),
    ("Boda", [("Ceremonia y recepción", "Variable"), ("Vestimenta", "Variable"), ("Proveedores", "Variable"),
              ("Otros gastos de boda", "Variable")]),
    ("Regalos", [("Regalos", "Variable"), ("Celebraciones", "Variable")]),
    ("Imprevistos", [("Emergencias", "Variable"), ("Multas y penalidades", "Variable"), ("Otros imprevistos", "Variable")]),
    ("Otros gastos", [("Comisiones bancarias", "Variable"), ("Donaciones", "Variable"), ("Otros", "Variable")]),
    ("Ingresos", [("Salario", "Fijo"), ("Bonificaciones", "Variable"), ("Ingresos extra", "Variable"),
                  ("Intereses y rendimientos", "Variable"), ("Otros ingresos", "Variable")]),
]
r = CAT0
for cat, subs in CATS:
    for s in subs:
        sub, clas = s[0], s[1]
        grupo = s[2] if len(s) > 2 else None
        comida = s[3] if len(s) > 3 else None
        tipo = "Ingreso" if cat == "Ingresos" else "Gasto"
        for j, v in enumerate([cat, sub, tipo, clas, grupo, comida]):
            put(ws, f"{L(2 + j)}{r}", v, F(10, j == 0, INK if j == 0 else TEXT), CARD, border=BOX,
                align=LEFT if j in (0, 1, 4) else CENTER)
        r += 1
for rr in range(r, CATL + 1):
    for j in range(6):
        put(ws, f"{L(2 + j)}{rr}", None, F(10), CARD, border=BOX, align=LEFT if j in (0, 1, 4) else CENTER)
for rr in range(CAT0, CATL + 1):
    ws[f"P{rr}"] = f'=IF(OR(B{rr}="",C{rr}=""),"",B{rr}&"|"&C{rr})'
    ws[f"Q{rr}"] = f'=IF(B{rr}="","",IF(COUNTIF(B${CAT0}:B{rr},B{rr})=1,MAX(Q${CAT0 - 1}:Q{rr - 1})+1,""))'
    ws[f"R{rr}"] = f'=IF(OR(F{rr}="",C{rr}=""),"",COUNTIF(F${CAT0}:F{rr},"?*"))'
for k in range(40):
    rr = CAT0 + k
    put(ws, f"L{rr}", f'=IFERROR(INDEX(catCat,MATCH({k + 1},catIdx,0))&"","")', F(10, True, INK), SOFT, border=BOX, align=LEFT)
    put(ws, f"M{rr}", f'=IF(L{rr}="","",IFERROR(INDEX(catTipo,MATCH(L{rr},catCat,0))&"",""))', F(9, False, MUTED), SOFT, border=BOX, align=CENTER)
ws.freeze_panes = "B6"
dv = DataValidation(type="list", formula1="=lstTipoCat", allow_blank=True)
dv.add(f"D{CAT0}:D{CATL}")
ws.add_data_validation(dv)
dv = DataValidation(type="list", formula1="=lstClasif", allow_blank=True)
dv.add(f"E{CAT0}:E{CATL}")
ws.add_data_validation(dv)
dv = DataValidation(type="list", formula1="=lstGrupos", allow_blank=True)
dv.add(f"F{CAT0}:F{CATL}")
ws.add_data_validation(dv)
dv = DataValidation(type="list", formula1="=lstComidas", allow_blank=True)
dv.add(f"G{CAT0}:G{CATL}")
ws.add_data_validation(dv)
ws.conditional_formatting.add(f"B{CAT0}:G{CATL}", FormulaRule(
    formula=[f'AND($B{CAT0}<>"",$B{CAT0}<>$B{CAT0 - 1},COUNTIF($B${CAT0 - 1}:$B{CAT0 - 1},$B{CAT0})>0)'],
    fill=fill(BAD_BG)))
note(ws, "L48", "Fila en rojo = categoría separada de su grupo: ordena la tabla por Categoría.")

# ============================================================================
# CUENTAS
# ============================================================================
ws = CTA
header(ws, "Cuentas y métodos de pago", "Saldos, conciliación y gasto por método", 21)
widths = {"A": 2, "B": 24, "C": 18, "D": 9, "E": 20, "F": 14, "G": 12, "H": 3, "I": 9, "J": 15, "K": 17,
          "L": 16, "M": 16, "N": 16, "O": 9, "P": 3, "Q": 15, "R": 12, "S": 16, "T": 15, "U": 2}
for col, w in widths.items():
    ws.column_dimensions[col].width = w
merge(ws, "B3:T4",
      "Registra cada fuente de dinero. «Vinculada a» se usa en tarjetas de débito: su gasto se reporta por separado, pero "
      "el saldo se descuenta de la cuenta bancaria vinculada. Las tarjetas de crédito muestran saldo negativo = deuda. "
      "El saldo solo se calcula si escribes un SALDO INICIAL real (cuenta movimientos desde su fecha, inclusive). "
      "Para conciliar: escribe el saldo que muestra tu banco y su fecha; si hay diferencia, registra un movimiento tipo Ajuste.",
      font=F(9, False, MUTED, True), align=TOP_WRAP)
ws.row_dimensions[3].height = 22
ws.row_dimensions[4].height = 22
merge(ws, "B5:G5", "DATOS DE LA CUENTA (editable)", font=F(9, True, "1E3A8A"), fill_c="E8EEF6", align=LEFT)
merge(ws, "I5:O5", "CÁLCULO AUTOMÁTICO", font=F(9, True, MUTED), fill_c="EEF2F7", align=LEFT)
merge(ws, "Q5:T5", "CONCILIACIÓN (opcional)", font=F(9, True, "1E3A8A"), fill_c="E8EEF6", align=LEFT)
cta_heads = {"B": "Cuenta / método", "C": "Tipo", "D": "Moneda", "E": "Vinculada a", "F": "Saldo inicial",
             "G": "Fecha saldo", "I": "¿Liquidez?", "J": "Movimiento neto", "K": "Saldo actual",
             "L": "Saldo en DOP", "M": "Gasto del período", "N": "Gasto mes actual", "O": "Nº movs",
             "Q": "Saldo según banco", "R": "Fecha de corte", "S": "Saldo calculado al corte", "T": "Diferencia"}
for col, h in cta_heads.items():
    put(ws, f"{col}6", h, F(9, True, "FFFFFF"), INK if col in "BCDEFGQR" else "475569", align=CENTER)
ws.row_dimensions[6].height = 30
ACCOUNTS = [
    ("Efectivo", "Efectivo", "DOP", None),
    ("Banco principal", "Cuenta bancaria", "DOP", None),
    ("Tarjeta de débito", "Tarjeta de débito", "DOP", "Banco principal"),
    ("Tarjeta de crédito", "Tarjeta de crédito", "DOP", None),
    ("Billetera digital", "Billetera digital", "DOP", None),
    ("Cuenta de trading", "Cuenta de trading", "USD", None),
]
for i in range(CTAN):
    r = CTA0 + i
    acc = ACCOUNTS[i] if i < len(ACCOUNTS) else (None, None, None, None)
    for col, v in zip("BCDE", acc):
        put(ws, f"{col}{r}", v, F(10, col == "B", INK), INPUT_BG, border=BOX, align=LEFT if col in "BE" else CENTER)
    put(ws, f"F{r}", None, F(10, True, "1E3A8A"), INPUT_BG, MONEY2, RIGHT, BOX)
    put(ws, f"G{r}", None, F(10, False, "1E3A8A"), INPUT_BG, DATE, CENTER, BOX)
    put(ws, f"Q{r}", None, F(10, True, "1E3A8A"), INPUT_BG, MONEY2, RIGHT, BOX)
    put(ws, f"R{r}", None, F(10, False, "1E3A8A"), INPUT_BG, DATE, CENTER, BOX)
    crit_ini = f'">="&IF(G{r}="",0,G{r})'
    f_mov = (f'IF(B{r}="","",SUMIFS(rEfO,rCtaSO,B{r},rFecha,{crit_ini})+SUMIFS(rEfD,rCtaSD,B{r},rFecha,{crit_ini}))')
    cells = {
        "I": f'=IF(B{r}="","",IF(COUNTIFS(tcTipo,C{r},tcLiq,"Sí")>0,"Sí","No"))',
        "J": "=" + f_mov,
        "K": f'=IF(B{r}="","",IF(E{r}<>"","↳ ver "&E{r},IF(F{r}="","Sin saldo inicial",F{r}+J{r})))',
        "L": f'=IF(ISNUMBER(K{r}),K{r}*IF(D{r}="USD",TasaUSD,1),"")',
        "M": f'=IF(B{r}="","",SUMIFS(rGasto,rCtaE,B{r},rFecha,">="&pIni,rFecha,"<="&pFin))',
        "N": f'=IF(B{r}="","",SUMIFS(rGasto,rCtaE,B{r},rFecha,">="&DATE(YEAR(TODAY()),MONTH(TODAY()),1),rFecha,"<="&EOMONTH(TODAY(),0)))',
        "O": f'=IF(B{r}="","",COUNTIF(rCtaE,B{r})+COUNTIF(rCtaSD,B{r}))',
        "S": f'=IF(OR(B{r}="",F{r}="",R{r}=""),"",F{r}+SUMIFS(rEfO,rCtaSO,B{r},rFecha,{crit_ini},rFecha,"<="&R{r})+SUMIFS(rEfD,rCtaSD,B{r},rFecha,{crit_ini},rFecha,"<="&R{r}))',
        "T": f'=IF(OR(Q{r}="",S{r}=""),"",Q{r}-S{r})',
    }
    for col, fml in cells.items():
        fmt = MONEY2 if col in "JKLMNST" else (NUM if col == "O" else None)
        put(ws, f"{col}{r}", fml, F(10, col in "KL", INK if col in "KL" else TEXT), SOFT, fmt,
            CENTER if col in "IO" else RIGHT, BOX)
r_tot = CTAL + 2
put(ws, f"B{r_tot}", "Totales en DOP", F(10, True, INK), align=LEFT)
put(ws, f"K{r_tot}", "Disponible (liquidez):", F(9, True, MUTED), align=RIGHT)
put(ws, f"L{r_tot}", '=SUMIFS(ctaSaldoDOP,ctaLiq,"Sí")', F(11, True, INK), fmt=MONEY2, align=RIGHT)
put(ws, f"K{r_tot + 1}", "Deuda tarjetas de crédito:", F(9, True, MUTED), align=RIGHT)
put(ws, f"L{r_tot + 1}", '=SUMIFS(ctaSaldoDOP,ctaTipo,"Tarjeta de crédito")', F(11, True, INK), fmt=MONEY2, align=RIGHT)
put(ws, f"K{r_tot + 2}", "Capital en trading:", F(9, True, MUTED), align=RIGHT)
put(ws, f"L{r_tot + 2}", '=SUMIFS(ctaSaldoDOP,ctaTipo,"Cuenta de trading")', F(11, True, INK), fmt=MONEY2, align=RIGHT)
put(ws, f"K{r_tot + 3}", "Cuentas con saldo inicial:", F(9, True, MUTED), align=RIGHT)
put(ws, f"L{r_tot + 3}", f'=SUMPRODUCT(--ISNUMBER(ctaSaldoDOP))&" de "&COUNTIFS(ctaNombre,"?*",ctaVinc,"")', F(11, True, INK), align=RIGHT)
put(ws, f"M{r_tot}", "Gasto total del período:", F(9, True, MUTED), align=RIGHT)
put(ws, f"N{r_tot}", f"=SUM(M{CTA0}:M{CTAL})", F(11, True, INK), fmt=MONEY2, align=RIGHT)
add_name(wb, "totDisponible", f"'Cuentas'!$L${r_tot}")
add_name(wb, "totDeudaTC", f"'Cuentas'!$L${r_tot + 1}")
add_name(wb, "totTrading", f"'Cuentas'!$L${r_tot + 2}")
add_name(wb, "nSaldosLiq", f"'Cuentas'!$O${r_tot}")
put(ws, f"O{r_tot}", '=SUMPRODUCT((ctaLiq="Sí")*ISNUMBER(ctaSaldoDOP))', F(8, False, "CBD5E1"))
put(ws, f"O{r_tot + 1}", '=COUNTIFS(ctaLiq,"Sí",ctaVinc,"")', F(8, False, "CBD5E1"))
add_name(wb, "nLiqTotal", f"'Cuentas'!$O${r_tot + 1}")
ws.freeze_panes = "C7"
for col, src in (("C", "=lstTiposCta"), ("D", "=lstMonedas"), ("E", "=lstCuentas")):
    dv = DataValidation(type="list", formula1=src, allow_blank=True)
    dv.add(f"{col}{CTA0}:{col}{CTAL}")
    ws.add_data_validation(dv)
dv = DataValidation(type="date", operator="between", formula1="DATE(2000,1,1)", formula2="DATE(2100,12,31)", allow_blank=True)
dv.add(f"G{CTA0}:G{CTAL}")
dv.add(f"R{CTA0}:R{CTAL}")
ws.add_data_validation(dv)
ws.conditional_formatting.add(f"T{CTA0}:T{CTAL}", FormulaRule(formula=[f'AND(ISNUMBER(T{CTA0}),ABS(T{CTA0})>=0.01)'], fill=fill(BAD_BG), font=Font(color=BAD, bold=True)))
ws.conditional_formatting.add(f"T{CTA0}:T{CTAL}", FormulaRule(formula=[f'AND(ISNUMBER(T{CTA0}),ABS(T{CTA0})<0.01)'], fill=fill(GOOD_BG), font=Font(color=GOOD, bold=True)))
ws.conditional_formatting.add(f"K{CTA0}:K{CTAL}", FormulaRule(formula=[f'ISTEXT(K{CTA0})'], font=Font(color="94A3B8", italic=True, bold=False)))
ws.conditional_formatting.add(f"M{CTA0}:M{CTAL}", DataBarRule(start_type="num", start_value=0, end_type="max", color="9EC5F4", showValue=True))

# ============================================================================
# REGISTRO
# ============================================================================
ws = REG
header(ws, "Registro de movimientos", "Una fila = una transacción", 35)
reg_widths = {"A": 2, "B": 11.5, "C": 15, "D": 24, "E": 28, "F": 34, "G": 13, "H": 8, "I": 9, "J": 19, "K": 19,
              "L": 22, "M": 11, "N": 17, "O": 30}
for col, w in reg_widths.items():
    ws.column_dimensions[col].width = w
for i in range(16, 36):
    ws.column_dimensions[L(i)].width = 14
ws.row_dimensions[3].height = 26
put(ws, "B3", f'=HYPERLINK("#\'Registro\'!B"&IFERROR(MATCH(TRUE,INDEX(rFecha="",0),0)+{R0 - 1},{RL}),"＋ Nuevo movimiento")',
    F(10, True, "FFFFFF"), BLUE, align=CENTER)
ws.merge_cells("B3:C3")
merge(ws, "D3:I3", "Atajos: Ctrl+; = fecha de hoy · Alt+↓ = abrir lista · Tab = siguiente campo · "
      "Solo Fecha y Monto son obligatorios: si dejas Tipo vacío se asume «Gasto».",
      font=F(9, False, MUTED, True), align=LEFT_WRAP)
put(ws, "J3", '="Movimientos: "&COUNTIF(rTipoE,"?*")', F(9, True, INK), align=LEFT)
put(ws, "K3", '="Advertencias: "&COUNTIF(rEstado,"⚠*")', F(9, True, BAD), align=LEFT)
put(ws, "L3", '=IF(COUNTIF(rObs,"DEMOSTRACIÓN")>0,"Filas DEMO: "&COUNTIF(rObs,"DEMOSTRACIÓN"),"")', F(9, True, VIOLET), align=LEFT)
merge(ws, "B4:N4", "TUS DATOS · llena estas columnas", font=F(9, True, "1E3A8A"), fill_c="E8EEF6", align=LEFT)
merge(ws, "O4:AI4", "CÁLCULOS AUTOMÁTICOS · no editar (pulsa «+» arriba de la columna P para ver el detalle)",
      font=F(9, True, MUTED), fill_c="EEF2F7", align=LEFT)
reg_heads = {
    "B": "Fecha", "C": "Tipo", "D": "Categoría", "E": "Subcategoría", "F": "Descripción", "G": "Monto",
    "H": "Moneda", "I": "Tasa USD", "J": "Cuenta / método", "K": "Cuenta destino", "L": "Comercio",
    "M": "Fijo/Variable", "N": "Observaciones", "O": "Estado", "P": "Tipo efectivo", "Q": "Cuenta efectiva",
    "R": "Tasa aplicada", "S": "Monto DOP", "T": "Monto USD", "U": "Gasto real DOP", "V": "Ingreso DOP",
    "W": "Resultado trading DOP", "X": "Clasificación efectiva", "Y": "Grupo alimentación", "Z": "Día semana",
    "AA": "Es hormiga", "AB": "Cuenta saldo origen", "AC": "Cuenta saldo destino", "AD": "Tipo cta origen",
    "AE": "Tipo cta destino", "AF": "Efecto saldo origen", "AG": "Efecto saldo destino", "AH": "Flujo líquido DOP",
    "AI": "Índice comercio",
}
for col, h in reg_heads.items():
    is_input = col in "BCDEFGHIJKLMN" and len(col) == 1
    put(ws, f"{col}5", h, F(10, True, "FFFFFF"), INK if is_input else "64748B", align=CENTER)
ws.row_dimensions[5].height = 30
tips = {
    "B": "Fecha de la transacción (Ctrl+; = hoy). Las compras con tarjeta de crédito van en la fecha de CONSUMO.",
    "C": "Vacío = Gasto. Transferencia / Pago tarjeta requieren Cuenta destino. Ver hoja Config para el detalle de cada tipo.",
    "E": "La lista depende de la Categoría elegida.",
    "G": "Siempre positivo, excepto Ajuste y Resultado trading (+ ganancia / − pérdida).",
    "H": "Vacío = DOP.",
    "I": "Opcional. Tasa real aplicada si el movimiento fue en USD (o entre cuentas de distinta moneda). Vacío = tasa de Config.",
    "J": "Cuenta o método con el que pagaste / donde entró el dinero.",
    "K": "Solo para Transferencia y Pago tarjeta.",
    "L": "Opcional. La lista sugiere comercios ya usados; puedes escribir uno nuevo.",
    "M": "Opcional: si queda vacío se usa la clasificación sugerida de la categoría.",
}
for col, t in tips.items():
    ws[f"{col}5"].comment = Comment(t, "Sistema")

REG_TOTAL_COLS = 35  # hasta AI


def reg_formulas(r):
    """Fórmulas de las columnas automáticas de la fila r."""
    p = f"P{r}"
    out = {}
    out["O"] = (f'=IF(AND(B{r}="",G{r}="",C{r}="",D{r}="",J{r}=""),"",'
                f'IF(B{r}="","⚠ Falta fecha",IF(G{r}="","⚠ Falta monto",IF(Q{r}="","⚠ Falta cuenta",'
                f'IF(ISNA(MATCH(Q{r},ctaNombre,0)),"⚠ Cuenta no registrada",'
                f'IF(AND(OR({p}="Gasto",{p}="Reembolso"),D{r}=""),"⚠ Falta categoría",'
                f'IF(AND(D{r}<>"",E{r}<>"",ISNA(MATCH(D{r}&"|"&E{r},catClave,0))),"⚠ Subcategoría no corresponde",'
                f'IF(AND(OR({p}="Transferencia",{p}="Pago tarjeta"),K{r}=""),"⚠ Falta cuenta destino",'
                f'IF(AND(K{r}<>"",{p}<>"Transferencia",{p}<>"Pago tarjeta"),"⚠ Destino solo aplica a transferencias",'
                f'IF(AND(K{r}<>"",K{r}=Q{r}),"⚠ Origen igual a destino",'
                f'IF(AND(K{r}<>"",ISNA(MATCH(K{r},ctaNombre,0))),"⚠ Destino no registrado",'
                f'IF(AND({p}="Pago tarjeta",AE{r}<>"Tarjeta de crédito"),"⚠ El destino debe ser tarjeta de crédito",'
                f'IF(AND({p}="Gasto",E{r}="Depósitos y retiros"),"⚠ Usa tipo Transferencia",'
                f'IF(AND({p}="Resultado trading",AD{r}<>"Cuenta de trading"),"⚠ Usa una cuenta de trading",'
                f'IF(AND(G{r}<=0,{p}<>"Ajuste",{p}<>"Resultado trading"),"⚠ El monto debe ser positivo","✓")))))))))))))))')
    out["P"] = f'=IF(OR(B{r}="",G{r}=""),"",IF(C{r}="","Gasto",C{r}))'
    out["Q"] = f'=IF({p}="","",IF(J{r}<>"",J{r},CuentaPred&""))'
    out["R"] = f'=IF({p}="","",IF(ISNUMBER(I{r}),I{r},TasaUSD))'
    out["S"] = f'=IF({p}="","",IF(H{r}="USD",G{r}*R{r},G{r}))'
    out["T"] = f'=IF({p}="","",IF(H{r}="USD",G{r},IF(R{r}>0,G{r}/R{r},0)))'
    out["U"] = f'=IF({p}="","",IF({p}="Gasto",S{r},IF({p}="Reembolso",-S{r},0)))'
    out["V"] = f'=IF({p}="","",IF({p}="Ingreso",S{r},0))'
    out["W"] = f'=IF({p}="","",IF({p}="Resultado trading",S{r},0))'
    out["X"] = (f'=IF(OR({p}="Gasto",{p}="Reembolso"),IF(M{r}<>"",M{r},'
                f'IFERROR(INDEX(catClasif,MATCH(D{r}&"|"&E{r},catClave,0))&"","Variable")),"")')
    out["Y"] = f'=IF({p}="","",IFERROR(INDEX(catGrupo,MATCH(D{r}&"|"&E{r},catClave,0))&"",""))'
    out["Z"] = f'=IF({p}="","",WEEKDAY(B{r},2))'
    out["AA"] = f'=IF({p}="","",IF(AND({p}="Gasto",X{r}="Variable",D{r}<>catTrading,S{r}>0,S{r}<=Umbral),1,0))'
    out["AB"] = (f'=IF(OR({p}="",Q{r}=""),"",IFERROR(IF(INDEX(ctaVinc,MATCH(Q{r},ctaNombre,0))&""<>"",'
                 f'INDEX(ctaVinc,MATCH(Q{r},ctaNombre,0))&"",Q{r}),Q{r}))')
    out["AC"] = (f'=IF(OR(K{r}="",AND({p}<>"Transferencia",{p}<>"Pago tarjeta")),"",'
                 f'IFERROR(IF(INDEX(ctaVinc,MATCH(K{r},ctaNombre,0))&""<>"",INDEX(ctaVinc,MATCH(K{r},ctaNombre,0))&"",K{r}),K{r}))')
    out["AD"] = f'=IF(AB{r}="","",IFERROR(INDEX(ctaTipo,MATCH(AB{r},ctaNombre,0))&"",""))'
    out["AE"] = f'=IF(AC{r}="","",IFERROR(INDEX(ctaTipo,MATCH(AC{r},ctaNombre,0))&"",""))'
    sign = f'IF(OR({p}="Gasto",{p}="Transferencia",{p}="Pago tarjeta"),-1,1)'
    out["AF"] = (f'=IF(OR({p}="",AB{r}=""),"",{sign}*IF(IFERROR(INDEX(ctaMoneda,MATCH(AB{r},ctaNombre,0))&"","DOP")="USD",T{r},S{r}))')
    out["AG"] = f'=IF(AC{r}="","",IF(IFERROR(INDEX(ctaMoneda,MATCH(AC{r},ctaNombre,0))&"","DOP")="USD",T{r},S{r}))'
    out["AH"] = (f'=IF({p}="","",IF(AND(AD{r}<>"",COUNTIFS(tcTipo,AD{r},tcLiq,"Sí")>0),{sign}*S{r},0)'
                 f'+IF(AND(AE{r}<>"",COUNTIFS(tcTipo,AE{r},tcLiq,"Sí")>0),S{r},0))')
    out["AI"] = (f'=IF(OR(L{r}="",{p}=""),"",IF(COUNTIFS(L${R0}:L{r},L{r},P${R0}:P{r},"?*")=1,'
                 f'MAX(AI${R0 - 1}:AI{r - 1})+1,""))')
    return out


# --- Datos DEMO -------------------------------------------------------------
def build_demo():
    rnd = random.Random(2026)
    rows = []
    start, end = date(2026, 8, 1), date(2026, 10, 14)

    def add(d, tipo, cat, sub, desc, monto, cuenta, mon="DOP", dest=None, com=None, clas=None, tasa=None):
        if d > end:
            return
        rows.append((d, tipo, cat, sub, "DEMO · " + desc, round(monto, 2), mon, tasa, cuenta, dest, com, clas))

    d = start
    while d <= end:
        wd = d.weekday()  # 0=lun
        if wd < 5:
            if rnd.random() < 0.4:
                add(d, "Gasto", "Alimentación", "Desayunos comprados", "Desayuno en la oficina", rnd.randrange(180, 340, 10),
                    rnd.choice(["Tarjeta de débito", "Efectivo"]), com="Cafetería del trabajo")
            if rnd.random() < 0.7:
                add(d, "Gasto", "Alimentación", "Almuerzos laborales", "Almuerzo laboral", rnd.randrange(350, 600, 25),
                    rnd.choice(["Tarjeta de débito", "Billetera digital", "Tarjeta de débito"]), com="Comedor ejecutivo")
            if rnd.random() < 0.4:
                add(d, "Gasto", "Alimentación", "Bebidas, snacks y otros consumos", rnd.choice(["Café", "Snack de la tarde", "Jugo y galletas"]),
                    rnd.randrange(90, 230, 10), rnd.choice(["Efectivo", "Billetera digital"]),
                    com=rnd.choice(["Cafetería del trabajo", "Colmado del barrio"]))
            if wd in (1, 3) and rnd.random() < 0.4:
                add(d, "Gasto", "Transporte y vehículo", "Parqueos y peajes", "Parqueo", rnd.choice([50, 75, 100]), "Efectivo", com="Parqueo centro")
        if rnd.random() < 0.15:
            add(d, "Gasto", "Alimentación", "Cenas compradas", "Cena comprada", rnd.randrange(400, 750, 25),
                rnd.choice(["Tarjeta de crédito", "Billetera digital"]), com=rnd.choice(["Pica pollo de la esquina", "Pizzería del sector"]))
        if wd == 4 and rnd.random() < 0.7:
            add(d, "Gasto", "Alimentación", "Delivery", "Pedido a domicilio", rnd.randrange(650, 1250, 50), "Tarjeta de crédito", com="App de delivery")
        if wd == 5 and d.isocalendar()[1] % 2 == 0:
            add(d, "Gasto", "Alimentación", "Restaurantes", "Almuerzo de fin de semana", rnd.randrange(1800, 3600, 100), "Tarjeta de crédito", com="Restaurante del malecón")
        if d.day in (2, 12, 22):
            add(d, "Gasto", "Alimentación", "Supermercado", "Compra de supermercado", rnd.randrange(3200, 6500, 50), "Tarjeta de crédito", com="Supermercado principal")
        if wd == 6 and rnd.random() < 0.7:
            add(d, "Gasto", "Alimentación", "Ingredientes para cocinar", "Vegetales, carne y víveres", rnd.randrange(350, 950, 25), "Efectivo", com="Mercado del barrio")
        if wd == 0:
            add(d, "Gasto", "Transporte y vehículo", "Combustible", "Combustible", rnd.randrange(2000, 3100, 50), "Tarjeta de crédito", com="Estación de combustible")
        d += timedelta(days=1)

    for m in (8, 9, 10):
        add(date(2026, m, 3), "Gasto", "Tecnología y suscripciones", "Herramientas de inteligencia artificial", "Suscripción herramienta IA", 20, "Tarjeta de crédito", "USD", com="Herramienta IA")
        add(date(2026, m, 5), "Gasto", "Hogar y servicios", "Internet", "Internet del hogar", 2050, "Banco principal", com="Proveedor de internet")
        if m < 10 or True:
            add(date(2026, m, 8), "Gasto", "Salud", "Gimnasio y bienestar", "Mensualidad gimnasio", 2500, "Tarjeta de crédito", com="Gimnasio")
        add(date(2026, m, 12), "Gasto", "Hogar y servicios", "Electricidad", "Factura de luz", rnd.randrange(2700, 3400, 10), "Banco principal", com="Empresa eléctrica")
        add(date(2026, m, 13), "Gasto", "Hogar y servicios", "Agua", "Factura de agua", 450, "Banco principal", com="Acueducto")
        add(date(2026, m, 15), "Ingreso", "Ingresos", "Salario", "Salario quincenal (ficticio)", 55000, "Banco principal")
        if m < 10:
            add(date(2026, m, 18), "Gasto", "Tecnología y suscripciones", "Entretenimiento", "Streaming de video", 10.99, "Tarjeta de crédito", "USD", com="Streaming")
            add(date(2026, m, 21), "Gasto", "Trading y actividades financieras", "Plataformas y herramientas", "Plataforma de gráficos", 39, "Tarjeta de crédito", "USD", com="Plataforma de trading")
            add(date(2026, m, 25), "Pago tarjeta", None, None, "Pago de tarjeta de crédito", 30000, "Banco principal", dest="Tarjeta de crédito")
            add(date(2026, m, 30), "Ingreso", "Ingresos", "Salario", "Salario quincenal (ficticio)", 55000, "Banco principal")
            add(date(2026, m, 16), "Transferencia", None, None, "Recarga de billetera digital", 4000, "Banco principal", dest="Billetera digital")
        add(date(2026, m, 1), "Transferencia", None, None, "Retiro en cajero", 5000, "Banco principal", dest="Efectivo")
        add(date(2026, m, 1) + timedelta(days=15) if m < 10 else date(2026, 10, 10), "Transferencia", None, None, "Retiro en cajero", 4000, "Banco principal", dest="Efectivo")
        add(date(2026, m, 10), "Gasto", "Pareja y actividades sociales", "Salidas y citas", "Salida con pareja", rnd.randrange(1800, 3200, 100), "Tarjeta de crédito", com="Cine y cena")
    # Trading: depósito (transferencia), evaluación (gasto), comisiones, resultados realizados y retiro
    add(date(2026, 8, 6), "Transferencia", "Trading y actividades financieras", "Depósitos y retiros", "Depósito a cuenta de trading", 15750, "Banco principal", dest="Cuenta de trading", tasa=63.0)
    add(date(2026, 8, 7), "Gasto", "Trading y actividades financieras", "Evaluaciones y cuentas de trading", "Evaluación de cuenta fondeada", 150, "Tarjeta de crédito", "USD", com="Firma de evaluación")
    add(date(2026, 8, 20), "Resultado trading", "Trading y actividades financieras", "Ganancias y pérdidas realizadas", "Pérdida realizada", -85, "Cuenta de trading", "USD")
    add(date(2026, 8, 20), "Gasto", "Trading y actividades financieras", "Comisiones", "Comisiones del broker", 4.5, "Cuenta de trading", "USD", com="Broker")
    add(date(2026, 9, 10), "Resultado trading", "Trading y actividades financieras", "Ganancias y pérdidas realizadas", "Ganancia realizada", 140, "Cuenta de trading", "USD")
    add(date(2026, 9, 10), "Gasto", "Trading y actividades financieras", "Comisiones", "Comisiones del broker", 6, "Cuenta de trading", "USD", com="Broker")
    add(date(2026, 9, 28), "Resultado trading", "Trading y actividades financieras", "Ganancias y pérdidas realizadas", "Pérdida realizada", -60, "Cuenta de trading", "USD")
    add(date(2026, 10, 8), "Resultado trading", "Trading y actividades financieras", "Ganancias y pérdidas realizadas", "Ganancia realizada", 95, "Cuenta de trading", "USD")
    add(date(2026, 10, 9), "Transferencia", "Trading y actividades financieras", "Depósitos y retiros", "Retiro desde cuenta de trading", 100, "Cuenta de trading", "USD", dest="Banco principal", tasa=63.2)
    # Otros
    add(date(2026, 8, 14), "Gasto", "Compras personales", "Ropa y calzado", "Camisa y pantalón", 3200, "Tarjeta de crédito", com="Tienda de ropa")
    add(date(2026, 8, 27), "Gasto", "Salud", "Medicamentos y farmacia", "Farmacia", 650, "Tarjeta de débito", com="Farmacia")
    add(date(2026, 9, 2), "Gasto", "Transporte y vehículo", "Mantenimiento", "Cambio de aceite", 3500, "Tarjeta de crédito", com="Taller")
    add(date(2026, 9, 9), "Gasto", "Tecnología y suscripciones", "Software", "Licencia anual de software", 99, "Tarjeta de crédito", "USD", com="Software anual")
    add(date(2026, 9, 12), "Gasto", "Alimentación", "Restaurantes", "Cena de cumpleaños (pagué todo)", 4200, "Tarjeta de crédito", com="Restaurante del malecón")
    add(date(2026, 9, 14), "Reembolso", "Alimentación", "Restaurantes", "Amigo devolvió su parte de la cena", 1400, "Billetera digital")
    add(date(2026, 9, 20), "Gasto", "Boda", "Proveedores", "Abono a proveedor de la boda", 15000, "Banco principal", com="Proveedor boda")
    add(date(2026, 9, 24), "Gasto", "Regalos", "Regalos", "Regalo de cumpleaños", 2000, "Tarjeta de crédito", com="Tienda de regalos")
    add(date(2026, 9, 26), "Ingreso", "Ingresos", "Ingresos extra", "Trabajo independiente (ficticio)", 8000, "Billetera digital")
    add(date(2026, 10, 1), "Ajuste", None, None, "Diferencia al contar el efectivo", -150, "Efectivo")
    add(date(2026, 10, 4), "Gasto", "Imprevistos", "Emergencias", "Reparación de celular", 1800, "Efectivo", com="Taller de celulares")
    add(date(2026, 10, 6), "Gasto", "Educación", "Cursos y certificaciones", "Curso en línea", 15, "Tarjeta de crédito", "USD", com="Plataforma de cursos")
    add(date(2026, 10, 7), "Gasto", "Compras personales", "Cuidado personal", "Barbería y productos", 850, "Efectivo", com="Barbería")
    rows.sort(key=lambda x: x[0])
    return rows


DEMO = build_demo()
in_cols = "BCDEFGHIJKLMN"
for i in range(NROWS):
    r = R0 + i
    if i < len(DEMO):
        d, tipo, cat, sub, desc, monto, mon, tasa, cuenta, dest, com, clas = DEMO[i]
        vals = [d, tipo, cat, sub, desc, monto, mon, tasa, cuenta, dest, com, clas, "DEMOSTRACIÓN"]
        for col, v in zip(in_cols, vals):
            ws[f"{col}{r}"] = v
    for col, fml in reg_formulas(r).items():
        ws[f"{col}{r}"] = fml

# Formatos de columna (estilos por celda: rápido y estable)
f_in, f_calc = F(10, False, INK), F(9, False, MUTED)
for i in range(NROWS):
    r = R0 + i
    for col in in_cols:
        c = ws[f"{col}{r}"]
        c.font = f_in
    ws[f"B{r}"].number_format = DATE
    ws[f"G{r}"].number_format = "#,##0.00"
    ws[f"I{r}"].number_format = "0.00"
    ws[f"O{r}"].font = F(9, True, GOOD)
    for col in ("S", "T", "U", "V", "W", "AF", "AG", "AH"):
        ws[f"{col}{r}"].number_format = "#,##0.00"
    for col in ("P", "Q", "R", "S", "T", "U", "V", "W", "X", "Y", "Z", "AA", "AB", "AC", "AD", "AE", "AF", "AG", "AH", "AI"):
        ws[f"{col}{r}"].font = f_calc
ws.freeze_panes = f"C{R0}"

tbl = Table(displayName="tblMovimientos", ref=f"B5:AI{RL}")
tbl.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True, showColumnStripes=False)
ws.add_table(tbl)
ws.column_dimensions.group("P", "AI", hidden=True, outline_level=1)

# Validaciones
def dv_list(ws, src, rng, strict=True, prompt=None):
    d = DataValidation(type="list", formula1=src, allow_blank=True, showErrorMessage=strict)
    if strict:
        d.error = "Elige un valor de la lista."
        d.errorTitle = "Valor no válido"
    if prompt:
        d.prompt = prompt
        d.showInputMessage = False
    d.add(rng)
    ws.add_data_validation(d)


dvd = DataValidation(type="date", operator="between", formula1="DATE(2000,1,1)", formula2="DATE(2100,12,31)",
                     allow_blank=True, showErrorMessage=True, error="Escribe una fecha válida (dd/mm/aaaa).", errorTitle="Fecha")
dvd.add(f"B{R0}:B{RL}")
ws.add_data_validation(dvd)
dv_list(ws, "=lstTipos", f"C{R0}:C{RL}")
dv_list(ws, "=lstCat", f"D{R0}:D{RL}")
dv_list(ws, f"=OFFSET('Categorías'!$C${CAT0},MATCH($D{R0},catCat,0)-1,0,COUNTIF(catCat,$D{R0}),1)", f"E{R0}:E{RL}")
dvm = DataValidation(type="decimal", operator="between", formula1="-1000000000", formula2="1000000000",
                     allow_blank=True, showErrorMessage=True, error="Escribe un número.", errorTitle="Monto")
dvm.add(f"G{R0}:G{RL}")
ws.add_data_validation(dvm)
dv_list(ws, "=lstMonedas", f"H{R0}:H{RL}")
dvt = DataValidation(type="decimal", operator="greaterThan", formula1="0", allow_blank=True, showErrorMessage=True,
                     error="La tasa debe ser mayor que 0.", errorTitle="Tasa")
dvt.add(f"I{R0}:I{RL}")
ws.add_data_validation(dvt)
dv_list(ws, "=lstCuentas", f"J{R0}:K{RL}")
dv_list(ws, "=lstComercios", f"L{R0}:L{RL}", strict=False)
dv_list(ws, "=lstClasif", f"M{R0}:M{RL}")

# Formato condicional
ws.conditional_formatting.add(f"B{R0}:N{RL}", FormulaRule(formula=[f'$N{R0}="DEMOSTRACIÓN"'], fill=fill(DEMO_BG)))
ws.conditional_formatting.add(f"O{R0}:O{RL}", FormulaRule(formula=[f'LEFT($O{R0},1)="⚠"'], fill=fill(BAD_BG), font=Font(color=BAD, bold=True)))
ws.conditional_formatting.add(f"G{R0}:G{RL}", FormulaRule(formula=[f'OR($P{R0}="Ingreso",$P{R0}="Reembolso")'], font=Font(color=GOOD, bold=True)))
ws.conditional_formatting.add(f"G{R0}:G{RL}", FormulaRule(formula=[f'OR($P{R0}="Transferencia",$P{R0}="Pago tarjeta",$P{R0}="Ajuste")'], font=Font(color="64748B", italic=True)))
ws.conditional_formatting.add(f"G{R0}:G{RL}", FormulaRule(formula=[f'$P{R0}="Resultado trading"'], font=Font(color=VIOLET, bold=True)))

# ============================================================================
# MOTOR (cálculos intermedios)
# ============================================================================
ws = MOT
header(ws, "Motor de cálculo", "Hoja técnica: alimenta gráficos y tablas. No editar.", 110)
put(ws, "B3", "Esta hoja hace los cálculos intermedios (rankings, series mensuales y semanales). Puedes ocultarla.", F(9, False, MUTED, True))
ws.column_dimensions["A"].width = 2
ws.column_dimensions["B"].width = 26
ws.column_dimensions["C"].width = 16

# Parámetros del período
section(ws, "B5:C5", "Período activo", size=10)
mp = [
    ("diaHoy", "Hoy", "=TODAY()", DATE),
    ("minFecha", "Primera fecha", '=IF(COUNT(rFecha)=0,TODAY(),MIN(rFecha))', DATE),
    ("maxFecha", "Última fecha", '=IF(COUNT(rFecha)=0,TODAY(),MAX(rFecha))', DATE),
    ("mesSel", "Mes seleccionado", '=IFERROR(DATE(VALUE(LEFT(fMes,4)),VALUE(RIGHT(fMes,2)),1),DATE(YEAR(TODAY()),MONTH(TODAY()),1))', DATE),
    ("pK", "Índice período", '=IFERROR(MATCH(fPeriodo,lstPeriodos,0),1)', "0"),
    ("pIni", "Inicio", '=CHOOSE(pK,DATE(YEAR(diaHoy),MONTH(diaHoy),1),EOMONTH(diaHoy,-2)+1,mesSel,diaHoy-6,diaHoy-29,diaHoy-89,'
                       'DATE(YEAR(diaHoy),1,1),DATE(YEAR(diaHoy)-1,1,1),minFecha,IF(ISNUMBER(fDesde),fDesde,minFecha))', DATE),
    ("pFin", "Fin", '=CHOOSE(pK,EOMONTH(diaHoy,0),EOMONTH(diaHoy,-1),EOMONTH(mesSel,0),diaHoy,diaHoy,diaHoy,'
                     'DATE(YEAR(diaHoy),12,31),DATE(YEAR(diaHoy)-1,12,31),MAX(maxFecha,diaHoy),IF(ISNUMBER(fHasta),fHasta,MAX(maxFecha,diaHoy)))', DATE),
    ("pDias", "Días transcurridos del período", '=MAX(1,MIN(pFin,MAX(diaHoy,pIni))-pIni+1)', "0"),
    ("pFinEff", "Fin efectivo (para promedios)", '=pIni+pDias-1', DATE),
    ("critCat", "Criterio categoría", '=IF(OR(fCat="",fCat="Todas"),"<>@@@",fCat)', None),
    ("critCta", "Criterio cuenta", '=IF(OR(fCta="",fCta="Todas"),"<>@@@",fCta)', None),
    ("refMes", "Mes de referencia", '=DATE(YEAR(pFinEff),MONTH(pFinEff),1)', DATE),
    ("refSem", "Semana de referencia (lunes)", '=pFinEff-WEEKDAY(pFinEff,3)', DATE),
    ("calIni", "Inicio calendario (máx. 366 días)", '=MAX(pIni,pFinEff-365)', DATE),
    ("semHoy", "Lunes de esta semana", '=diaHoy-WEEKDAY(diaHoy,3)', DATE),
    ("mesHoy", "Inicio de este mes", '=DATE(YEAR(diaHoy),MONTH(diaHoy),1)', DATE),
    ("mesesDatos3", "Meses con datos (3 meses previos)", '=MAX(1,(COUNTIFS(rTipoE,"?*",rFecha,">="&EDATE(mesHoy,-3),rFecha,"<"&EDATE(mesHoy,-2))>0)'
                                                          '+(COUNTIFS(rTipoE,"?*",rFecha,">="&EDATE(mesHoy,-2),rFecha,"<"&EDATE(mesHoy,-1))>0)'
                                                          '+(COUNTIFS(rTipoE,"?*",rFecha,">="&EDATE(mesHoy,-1),rFecha,"<"&mesHoy)>0))', "0"),
    ("pTexto", "Texto del período", "=" + fecha_txt("pIni") + '&"  →  "&' + fecha_txt("pFin"), None),
]
for i, (nm, label, fml, fmt) in enumerate(mp):
    r = 6 + i
    put(ws, f"B{r}", label, F(9, False, TEXT))
    put(ws, f"C{r}", fml, F(9, True, INK), fmt=fmt)
    add_name(wb, nm, f"'Motor'!$C${r}")


def block_title(ws, col, text, width_cols):
    rng = f"{L(col)}5:{L(col + width_cols - 1)}5"
    section(ws, rng, text, size=10)


def hdr(ws, col, row, names):
    for j, h in enumerate(names):
        put(ws, f"{L(col + j)}{row}", h, F(9, True, "FFFFFF"), "475569", align=CENTER)
        ws.column_dimensions[L(col + j)].width = max(ws.column_dimensions[L(col + j)].width or 0, 13)


DATES = '">="&pIni'
DATEE = '"<="&pFin'
FILT_D = f"rFecha,{DATES},rFecha,{DATEE}"

# --- Bloque: categorías (sin filtro de categoría) -----------------------------
c0 = 5  # E
block_title(ws, c0, "Gasto por categoría", 5)
hdr(ws, c0, 6, ["k", "Categoría", "Gasto", "Clave orden", "Nº gastos"])
for k in range(1, 41):
    r = 6 + k
    E, Fc, G, H, I = (L(c0 + j) for j in range(5))
    ws[f"{E}{r}"] = k
    ws[f"{Fc}{r}"] = f'=IFERROR(INDEX(catUList,{k})&"","")'
    ws[f"{G}{r}"] = f'=IF({Fc}{r}="","",SUMIFS(rGasto,rCat,{Fc}{r},{FILT_D},rCtaE,critCta))'
    ws[f"{H}{r}"] = f'=IF({Fc}{r}="",-1E+12-{k},N({G}{r})+(100-{k})/1E+7)'
    ws[f"{I}{r}"] = f'=IF({Fc}{r}="","",COUNTIFS(rCat,{Fc}{r},rTipoE,"Gasto",{FILT_D},rCtaE,critCta))'
CATB = {"name": f"$F$7:$F$46", "val": f"$G$7:$G$46", "key": f"$H$7:$H$46", "cnt": f"$I$7:$I$46"}

# --- Bloque: subcategorías (con filtro de categoría) --------------------------
c0 = 11  # K
block_title(ws, c0, "Gasto por subcategoría (concepto)", 6)
hdr(ws, c0, 6, ["Categoría", "Subcategoría", "Gasto", "Nº gastos", "Clave orden", "Etiqueta"])
for k in range(1, CATN + 1):
    r = 6 + k
    K_, L_, M_, N_, O_, P_ = (L(c0 + j) for j in range(6))
    ws[f"{K_}{r}"] = f'=INDEX(catCat,{k})&""'
    ws[f"{L_}{r}"] = f'=INDEX(catSub,{k})&""'
    ws[f"{M_}{r}"] = (f'=IF(OR({K_}{r}="",{L_}{r}="",INDEX(catTipo,{k})&""<>"Gasto"),"",IF(AND(critCat<>"<>@@@",{K_}{r}<>critCat),0,'
                      f'SUMIFS(rGasto,rCat,{K_}{r},rSub,{L_}{r},{FILT_D},rCtaE,critCta)))')
    ws[f"{N_}{r}"] = (f'=IF({M_}{r}="","",IF(AND(critCat<>"<>@@@",{K_}{r}<>critCat),0,'
                      f'COUNTIFS(rCat,{K_}{r},rSub,{L_}{r},rTipoE,"Gasto",{FILT_D},rCtaE,critCta)))')
    ws[f"{O_}{r}"] = f'=IF({M_}{r}="",-1E+12-{k},N({M_}{r})+({CATN + 1}-{k})/1E+7)'
    ws[f"{P_}{r}"] = f'={L_}{r}&" · "&{K_}{r}'
SUBB = {"name": f"$P$7:$P${6 + CATN}", "val": f"$M$7:$M${6 + CATN}", "key": f"$O$7:$O${6 + CATN}", "cnt": f"$N$7:$N${6 + CATN}"}

# --- Bloque: comercios --------------------------------------------------------
c0 = 18  # R
NCOM = 300
block_title(ws, c0, "Comercios (lista automática)", 8)
hdr(ws, c0, 6, ["k", "Comercio", "Gasto", "Nº gastos", "Clave orden", "Gasto hormiga", "Nº hormiga", "Clave hormiga"])
for k in range(1, NCOM + 1):
    r = 6 + k
    R_, S_, T_, U_, V_, W_, X_, Y_ = (L(c0 + j) for j in range(8))
    ws[f"{R_}{r}"] = k
    ws[f"{S_}{r}"] = f'=IFERROR(INDEX(rCom,MATCH({k},rComIdx,0))&"","")'
    ws[f"{T_}{r}"] = f'=IF({S_}{r}="","",SUMIFS(rGasto,rCom,{S_}{r},{FILT_D},rCat,critCat,rCtaE,critCta))'
    ws[f"{U_}{r}"] = f'=IF({S_}{r}="","",COUNTIFS(rCom,{S_}{r},rTipoE,"Gasto",{FILT_D},rCat,critCat,rCtaE,critCta))'
    ws[f"{V_}{r}"] = f'=IF({S_}{r}="",-1E+12-{k},N({T_}{r})+({NCOM + 1}-{k})/1E+7)'
    ws[f"{W_}{r}"] = f'=IF({S_}{r}="","",SUMIFS(rGasto,rCom,{S_}{r},rHorm,1,{FILT_D}))'
    ws[f"{X_}{r}"] = f'=IF({S_}{r}="","",COUNTIFS(rCom,{S_}{r},rHorm,1,{FILT_D}))'
    ws[f"{Y_}{r}"] = f'=IF({S_}{r}="",-1E+12-{k},N({W_}{r})+({NCOM + 1}-{k})/1E+7)'
COMB = {"name": f"$S$7:$S${6 + NCOM}", "val": f"$T$7:$T${6 + NCOM}", "key": f"$V$7:$V${6 + NCOM}", "cnt": f"$U$7:$U${6 + NCOM}"}
COMH = {"name": f"$S$7:$S${6 + NCOM}", "val": f"$W$7:$W${6 + NCOM}", "key": f"$Y$7:$Y${6 + NCOM}", "cnt": f"$X$7:$X${6 + NCOM}"}
add_name(wb, "lstComercios", f"OFFSET('Motor'!$S$7,0,0,MAX(1,COUNTIF('Motor'!$S$7:$S${6 + NCOM},\"?*\")),1)")

# --- Bloque: cuentas / métodos ---------------------------------------------
c0 = 27  # AA
block_title(ws, c0, "Gasto por cuenta / método", 4)
hdr(ws, c0, 6, ["Cuenta", "Gasto", "Nº gastos", "Clave orden"])
for k in range(1, CTAN + 1):
    r = 6 + k
    A_, B_, C_, D_ = (L(c0 + j) for j in range(4))
    ws[f"{A_}{r}"] = f'=INDEX(ctaNombre,{k})&""'
    ws[f"{B_}{r}"] = f'=IF({A_}{r}="","",SUMIFS(rGasto,rCtaE,{A_}{r},{FILT_D},rCat,critCat))'
    ws[f"{C_}{r}"] = f'=IF({A_}{r}="","",COUNTIFS(rCtaE,{A_}{r},rTipoE,"Gasto",{FILT_D},rCat,critCat))'
    ws[f"{D_}{r}"] = f'=IF({A_}{r}="",-1E+12-{k},N({B_}{r})+(100-{k})/1E+7)'
CTAB = {"name": f"$AA$7:$AA${6 + CTAN}", "val": f"$AB$7:$AB${6 + CTAN}", "key": f"$AD$7:$AD${6 + CTAN}", "cnt": f"$AC$7:$AC${6 + CTAN}"}

# --- Bloque: hormiga por categoría ------------------------------------------
c0 = 32  # AF
block_title(ws, c0, "Gasto hormiga por categoría", 4)
hdr(ws, c0, 6, ["Categoría", "Gasto hormiga", "Nº", "Clave orden"])
for k in range(1, 41):
    r = 6 + k
    A_, B_, C_, D_ = (L(c0 + j) for j in range(4))
    ws[f"{A_}{r}"] = f'=$F{r}'
    ws[f"{B_}{r}"] = f'=IF({A_}{r}="","",SUMIFS(rGasto,rCat,{A_}{r},rHorm,1,{FILT_D}))'
    ws[f"{C_}{r}"] = f'=IF({A_}{r}="","",COUNTIFS(rCat,{A_}{r},rHorm,1,{FILT_D}))'
    ws[f"{D_}{r}"] = f'=IF({A_}{r}="",-1E+12-{k},N({B_}{r})+(100-{k})/1E+7)'
HORB = {"name": f"$AF$7:$AF$46", "val": f"$AG$7:$AG$46", "key": f"$AI$7:$AI$46", "cnt": f"$AH$7:$AH$46"}

# --- Bloque: 12 meses ---------------------------------------------------------
c0 = 37  # AK
mcols = ["Mes", "Etiqueta", "Ingresos", "Gasto real", "Ahorro neto", "Fijos", "Variables", "Comida preparada",
         "Súper e ingredientes", "Bebidas y snacks", "Hormiga", "Nº comidas compradas", "Resultado trading", "Tasa ahorro"]
block_title(ws, c0, "Últimos 12 meses (hasta el mes de referencia)", len(mcols))
hdr(ws, c0, 6, mcols)
MC = {h: L(c0 + j) for j, h in enumerate(mcols)}
for j in range(1, 13):
    r = 6 + j
    m = MC["Mes"]
    me = f'"<="&EOMONTH({m}{r},0)'
    ms = f'">="&{m}{r}'
    dd = f"rFecha,{ms},rFecha,{me}"
    ws[f"{m}{r}"] = f"=EDATE(refMes,{j - 12})"
    ws[f"{m}{r}"].number_format = "mmm yyyy"
    ws[f"{MC['Etiqueta']}{r}"] = f'=CHOOSE(MONTH({m}{r}),{MESES_ES})&" "&RIGHT(YEAR({m}{r}),2)'
    ws[f"{MC['Ingresos']}{r}"] = f"=SUMIFS(rIng,{dd},rCtaE,critCta)"
    ws[f"{MC['Gasto real']}{r}"] = f"=SUMIFS(rGasto,{dd},rCat,critCat,rCtaE,critCta)"
    ws[f"{MC['Ahorro neto']}{r}"] = f"={MC['Ingresos']}{r}-{MC['Gasto real']}{r}"
    ws[f"{MC['Fijos']}{r}"] = f'=SUMIFS(rGasto,{dd},rClasE,"Fijo",rCat,critCat,rCtaE,critCta)'
    ws[f"{MC['Variables']}{r}"] = f'=SUMIFS(rGasto,{dd},rClasE,"Variable",rCat,critCat,rCtaE,critCta)'
    ws[f"{MC['Comida preparada']}{r}"] = f"=SUMIFS(rGasto,{dd},rGrupo,gCP)"
    ws[f"{MC['Súper e ingredientes']}{r}"] = f"=SUMIFS(rGasto,{dd},rGrupo,gSI)"
    ws[f"{MC['Bebidas y snacks']}{r}"] = f"=SUMIFS(rGasto,{dd},rGrupo,gBS)"
    ws[f"{MC['Hormiga']}{r}"] = f"=SUMIFS(rGasto,{dd},rHorm,1)"
    ws[f"{MC['Nº comidas compradas']}{r}"] = f'=COUNTIFS({dd},rGrupo,gCP,rTipoE,"Gasto")'
    ws[f"{MC['Resultado trading']}{r}"] = f"=SUMIFS(rTrad,{dd})"
    ws[f"{MC['Tasa ahorro']}{r}"] = f'=IF({MC["Ingresos"]}{r}>0,{MC["Ahorro neto"]}{r}/{MC["Ingresos"]}{r},"")'
    for h in mcols[2:]:
        ws[f"{MC[h]}{r}"].number_format = PCT if h == "Tasa ahorro" else "#,##0"
ws.column_dimensions[MC["Mes"]].width = 12

# --- Bloque: 12 semanas -------------------------------------------------------
c0 = 52  # AZ
wcols = ["Semana (lunes)", "Etiqueta", "Gasto real", "Comida preparada", "Súper e ingredientes", "Hormiga"]
block_title(ws, c0, "Últimas 12 semanas", len(wcols))
hdr(ws, c0, 6, wcols)
WC = {h: L(c0 + j) for j, h in enumerate(wcols)}
for j in range(1, 13):
    r = 6 + j
    w0 = WC["Semana (lunes)"]
    dd = f'rFecha,">="&{w0}{r},rFecha,"<="&{w0}{r}+6'
    ws[f"{w0}{r}"] = f"=refSem-7*{12 - j}"
    ws[f"{w0}{r}"].number_format = DATE
    ws[f"{WC['Etiqueta']}{r}"] = f'=DAY({w0}{r})&" "&CHOOSE(MONTH({w0}{r}),{MESES_ES})'
    ws[f"{WC['Gasto real']}{r}"] = f"=SUMIFS(rGasto,{dd},rCat,critCat,rCtaE,critCta)"
    ws[f"{WC['Comida preparada']}{r}"] = f"=SUMIFS(rGasto,{dd},rGrupo,gCP)"
    ws[f"{WC['Súper e ingredientes']}{r}"] = f"=SUMIFS(rGasto,{dd},rGrupo,gSI)"
    ws[f"{WC['Hormiga']}{r}"] = f"=SUMIFS(rGasto,{dd},rHorm,1)"

# --- Bloque: calendario del período (alimentación por día) -------------------
c0 = 59  # BG
ccols = ["k", "Fecha", "Día semana", "Alimentación", "Comida preparada", "Clave orden"]
block_title(ws, c0, "Calendario del período (hasta 366 días)", len(ccols))
hdr(ws, c0, 6, ccols)
CC = {h: L(c0 + j) for j, h in enumerate(ccols)}
for k in range(1, 367):
    r = 6 + k
    f_ = CC["Fecha"]
    ws[f"{CC['k']}{r}"] = k
    ws[f"{f_}{r}"] = f'=IF(calIni+{k - 1}<=pFinEff,calIni+{k - 1},"")'
    ws[f"{f_}{r}"].number_format = DATE
    ws[f"{CC['Día semana']}{r}"] = f'=IF({f_}{r}="","",WEEKDAY({f_}{r},2))'
    ws[f"{CC['Alimentación']}{r}"] = f'=IF({f_}{r}="","",SUMIFS(rGasto,rFecha,{f_}{r},rGrupo,"?*"))'
    ws[f"{CC['Comida preparada']}{r}"] = f'=IF({f_}{r}="","",SUMIFS(rGasto,rFecha,{f_}{r},rGrupo,gCP))'
    ws[f"{CC['Clave orden']}{r}"] = f'=IF({f_}{r}="",-1E+12-{k},N({CC["Alimentación"]}{r})+(400-{k})/1E+7)'
CALB = {"date": f"${CC['Fecha']}$7:${CC['Fecha']}$372", "val": f"${CC['Alimentación']}$7:${CC['Alimentación']}$372",
        "key": f"${CC['Clave orden']}$7:${CC['Clave orden']}$372", "cp": f"${CC['Comida preparada']}$7:${CC['Comida preparada']}$372",
        "wd": f"${CC['Día semana']}$7:${CC['Día semana']}$372"}

# --- Bloque: alimentación por día de la semana --------------------------------
c0 = 66  # BN
dcols = ["Nº", "Día", "Alimentación", "Comida preparada", "Nº comidas compradas", "Días en período", "Promedio por día", "Comida prep. por día"]
block_title(ws, c0, "Alimentación por día de la semana", len(dcols))
hdr(ws, c0, 6, dcols)
DC = {h: L(c0 + j) for j, h in enumerate(dcols)}
for n, nombre in enumerate(["Lunes", "Martes", "Miércoles", "Jueves", "Viernes", "Sábado", "Domingo"], start=1):
    r = 6 + n
    ws[f"{DC['Nº']}{r}"] = n
    ws[f"{DC['Día']}{r}"] = nombre
    ws[f"{DC['Alimentación']}{r}"] = f'=SUMIFS(rGasto,rGrupo,"?*",rDia,{n},{FILT_D})'
    ws[f"{DC['Comida preparada']}{r}"] = f"=SUMIFS(rGasto,rGrupo,gCP,rDia,{n},{FILT_D})"
    ws[f"{DC['Nº comidas compradas']}{r}"] = f'=COUNTIFS(rGrupo,gCP,rTipoE,"Gasto",rDia,{n},{FILT_D})'
    ws[f"{DC['Días en período']}{r}"] = f"=MAX(0,INT((pFinEff-(pIni+MOD({n}-WEEKDAY(pIni,2),7)))/7)+1)"
    ws[f"{DC['Promedio por día']}{r}"] = f'=IF({DC["Días en período"]}{r}>0,{DC["Alimentación"]}{r}/{DC["Días en período"]}{r},0)'
    ws[f"{DC['Comida prep. por día']}{r}"] = f'=IF({DC["Días en período"]}{r}>0,{DC["Comida preparada"]}{r}/{DC["Días en período"]}{r},0)'
    for h in dcols[2:]:
        ws[f"{DC[h]}{r}"].number_format = "#,##0"


# --- Bloques ordenados (top N) -----------------------------------------------
def sorted_block(c0, title, src, n, sheet="Motor", extra_pct_total=None):
    """Top-N ordenado con LARGE + INDEX/MATCH (sin funciones dinámicas)."""
    block_title(ws, c0, title, 4)
    hdr(ws, c0, 6, ["#", "Nombre", "Monto", "Nº"])
    A_, B_, C_, D_ = (L(c0 + j) for j in range(4))
    for j in range(1, n + 1):
        r = 6 + j
        key = f"LARGE({src['key']},{j})"
        pos = f"MATCH({key},{src['key']},0)"
        ws[f"{A_}{r}"] = j
        ws[f"{B_}{r}"] = f'=IFERROR(IF(N(INDEX({src["val"]},{pos}))>0,INDEX({src["name"]},{pos})&"",""),"")'
        ws[f"{C_}{r}"] = f'=IF({B_}{r}="","",INDEX({src["val"]},{pos}))'
        ws[f"{D_}{r}"] = f'=IF({B_}{r}="","",INDEX({src["cnt"]},{pos}))'
        ws[f"{C_}{r}"].number_format = "#,##0"
    return {"name": f"'Motor'!${B_}$7:${B_}${6 + n}", "val": f"'Motor'!${C_}$7:${C_}${6 + n}",
            "cnt": f"'Motor'!${D_}$7:${D_}${6 + n}", "col_name": B_, "col_val": C_, "col_cnt": D_, "n": n}


S_CAT = sorted_block(75, "Top categorías", CATB, 12)
S_SUB = sorted_block(80, "Top 10 conceptos", SUBB, 10)
S_COM = sorted_block(85, "Top 10 comercios", COMB, 10)
S_CTA = sorted_block(90, "Top métodos de pago", CTAB, 8)
S_HCAT = sorted_block(95, "Top hormiga por categoría", HORB, 10)
S_HCOM = sorted_block(100, "Top hormiga por comercio", COMH, 10)
# Top días de alimentación
c0 = 105
block_title(ws, c0, "Días con más gasto en alimentación", 4)
hdr(ws, c0, 6, ["#", "Fecha", "Alimentación", "Comida preparada"])
for j in range(1, 8):
    r = 6 + j
    key = f"LARGE({CALB['key']},{j})"
    pos = f"MATCH({key},{CALB['key']},0)"
    A_, B_, C_, D_ = (L(c0 + i) for i in range(4))
    ws[f"{A_}{r}"] = j
    ws[f"{B_}{r}"] = f'=IFERROR(IF(N(INDEX({CALB["val"]},{pos}))>0,INDEX({CALB["date"]},{pos}),""),"")'
    ws[f"{B_}{r}"].number_format = DATE
    ws[f"{C_}{r}"] = f'=IF({B_}{r}="","",INDEX({CALB["val"]},{pos}))'
    ws[f"{D_}{r}"] = f'=IF({B_}{r}="","",INDEX({CALB["cp"]},{pos}))'
S_DAYS = {"date": L(106), "val": L(107), "cp": L(108)}
for c in range(5, 112):
    if ws.column_dimensions[L(c)].width in (None, 13):
        ws.column_dimensions[L(c)].width = 13
for row in ws.iter_rows(min_row=7, max_row=372, min_col=5, max_col=111):
    for c in row:
        if c.value is not None and c.font.name != FONT:
            c.font = F(9, False, TEXT)

# ============================================================================
# DASHBOARD
# ============================================================================
ws = DSH
LAST = 18
header(ws, "Dashboard financiero", "Todas las cifras en DOP · se actualiza solo", LAST)
ws.column_dimensions["A"].width = 2
for c in range(2, LAST):
    ws.column_dimensions[L(c)].width = 11
ws.column_dimensions[L(LAST)].width = 2
area(ws, f"A3:{L(LAST)}140", fill_c=PAGE)

# Filtros
labels = [("B", "D", "PERÍODO"), ("E", "G", "MES (si eliges «Mes seleccionado»)"), ("H", "I", "DESDE (personalizado)"),
          ("J", "K", "HASTA (personalizado)"), ("L", "N", "CATEGORÍA (aplica a gastos)"), ("O", "Q", "CUENTA / MÉTODO")]
ws.row_dimensions[4].height = 18
ws.row_dimensions[5].height = 26
for a, b, t in labels:
    merge(ws, f"{a}4:{b}4", t, font=F(8, True, MUTED), align=Alignment(horizontal="left", vertical="bottom"))
    ws.merge_cells(f"{a}5:{b}5")
    for row in ws[f"{a}5:{b}5"]:
        for c in row:
            c.fill = fill(CARD)
            c.border = INPUT_BOX
    ws[f"{a}5"].font = F(11, True, "1E3A8A")
    ws[f"{a}5"].alignment = CENTER
ws["B5"] = "Todo el historial"
ws["E5"] = "2026-10"
ws["H5"].number_format = DATE
ws["J5"].number_format = DATE
ws["L5"] = "Todas"
ws["O5"] = "Todas"
add_name(wb, "fPeriodo", "'Dashboard'!$B$5")
add_name(wb, "fMes", "'Dashboard'!$E$5")
add_name(wb, "fDesde", "'Dashboard'!$H$5")
add_name(wb, "fHasta", "'Dashboard'!$J$5")
add_name(wb, "fCat", "'Dashboard'!$L$5")
add_name(wb, "fCta", "'Dashboard'!$O$5")
for ref, src in (("B5", "=lstPeriodos"), ("E5", "=lstMeses"), ("L5", "=lstCatFiltro"), ("O5", "=lstCtaFiltro")):
    dv_list(ws, src, ref)
dv = DataValidation(type="date", operator="between", formula1="DATE(2000,1,1)", formula2="DATE(2100,12,31)", allow_blank=True)
dv.add("H5")
dv.add("J5")
ws.add_data_validation(dv)
merge(ws, "B6:Q6", '="Mostrando  "&pTexto&"   ·   "&pDias&" días transcurridos   ·   Categoría: "&IF(critCat="<>@@@","todas",fCat)&"   ·   Cuenta: "&IF(critCta="<>@@@","todas",fCta)',
      font=F(9, False, MUTED, True), align=Alignment(horizontal="left", vertical="center"))

GASTO_P = f"SUMIFS(rGasto,{FILT_D},rCat,critCat,rCtaE,critCta)"
ING_P = f"SUMIFS(rIng,{FILT_D},rCtaE,critCta)"


def kpi(ws, r, c1, c2, label, value, fmt, sub=None, accent=BLUE, value_size=20):
    a, b = L(c1), L(c2)
    area(ws, f"{a}{r}:{b}{r + 3}", fill_c=CARD)
    outline_box(ws, f"{a}{r}:{b}{r + 3}")
    for rr in range(r, r + 4):
        ws[f"{a}{rr}"].border = Border(left=Side(style="thick", color=accent),
                                       top=ws[f"{a}{rr}"].border.top, bottom=ws[f"{a}{rr}"].border.bottom)
    merge(ws, f"{a}{r}:{b}{r}", label.upper(), font=F(8, True, MUTED), align=Alignment(horizontal="left", vertical="bottom", indent=1))
    merge(ws, f"{a}{r + 1}:{b}{r + 2}", value, font=F(value_size, True, INK), fmt=fmt,
          align=Alignment(horizontal="left", vertical="center", indent=1))
    merge(ws, f"{a}{r + 3}:{b}{r + 3}", sub, font=F(8, False, MUTED), align=Alignment(horizontal="left", vertical="top", indent=1, wrap_text=True))


def kpi_rows(r):
    ws.row_dimensions[r].height = 18
    ws.row_dimensions[r + 1].height = 18
    ws.row_dimensions[r + 2].height = 18
    ws.row_dimensions[r + 3].height = 24


kpi_rows(8)
kpi(ws, 8, 2, 5, "Ingresos del período", f"={ING_P}", MONEY,
    '="Promedio mensual (meses con ingresos): RD$ "&FIXED(IFERROR(AVERAGEIF(\'Motor\'!$AM$7:$AM$18,"<>0"),0),0)', AQUA)
kpi(ws, 8, 6, 9, "Gasto real del período", f"={GASTO_P}", MONEY,
    f'="Bruto RD$ "&FIXED(SUMIFS(rDOP,rTipoE,"Gasto",{FILT_D},rCat,critCat,rCtaE,critCta),0)&" − reembolsos RD$ "&FIXED(SUMIFS(rDOP,rTipoE,"Reembolso",{FILT_D},rCat,critCat,rCtaE,critCta),0)', ORANGE)
kpi(ws, 8, 10, 13, "Ahorro neto (ingresos − gasto)", f"={ING_P}-{GASTO_P}", MONEY,
    f'=IF({ING_P}>0,"Tasa de ahorro: "&FIXED(({ING_P}-{GASTO_P})/{ING_P}*100,1)&"%   ·   sin incluir trading","Sin ingresos en el período")', BLUE)
kpi(ws, 8, 14, 17, "Dinero disponible", '=IF(nSaldosLiq=0,"Sin saldos",totDisponible)', MONEY,
    '=IF(nSaldosLiq=0,"Registra saldos iniciales en Cuentas para calcularlo",nSaldosLiq&" de "&nLiqTotal&" cuentas líquidas  ·  Deuda TC: RD$ "&FIXED(-totDeudaTC,0)&"  ·  Trading: RD$ "&FIXED(totTrading,0))', VIOLET)
kpi_rows(13)
FCC = "rCat,critCat,rCtaE,critCta"
kpi(ws, 13, 2, 5, "Gasto de hoy", f'=SUMIFS(rGasto,rFecha,diaHoy,{FCC})', MONEY,
    f'="Movimientos diaHoy: "&COUNTIFS(rFecha,diaHoy,rTipoE,"Gasto",{FCC})', ORANGE, 16)
kpi(ws, 13, 6, 9, "Gasto de esta semana", f'=SUMIFS(rGasto,rFecha,">="&semHoy,rFecha,"<="&semHoy+6,{FCC})', MONEY,
    '="Lunes "&' + fecha_txt("semHoy") + '&" a domingo"', ORANGE, 16)
kpi(ws, 13, 10, 13, "Gasto de este mes", f'=SUMIFS(rGasto,rFecha,">="&mesHoy,rFecha,"<="&EOMONTH(diaHoy,0),{FCC})', MONEY,
    '="Mes de "&CHOOSE(MONTH(diaHoy),' + MESES_ES + ')&" "&YEAR(diaHoy)', ORANGE, 16)
kpi(ws, 13, 14, 17, "Promedio diario del período", f"={GASTO_P}/pDias", MONEY,
    f'="Promedio semanal: RD$ "&FIXED({GASTO_P}/pDias*7,0)&"  ·  mensual: RD$ "&FIXED({GASTO_P}/pDias*30.4,0)', BLUE, 16)

# Tres lentes: flujo, gasto real, trading
section(ws, "B18:Q18", "Tres lentes distintos: flujo de efectivo · gasto real · trading")
ws.row_dimensions[18].height = 24


def lens(c1, c2, title, rows, color, footnote):
    a, b = L(c1), L(c2)
    area(ws, f"{a}19:{b}25", fill_c=CARD)
    outline_box(ws, f"{a}19:{b}25")
    merge(ws, f"{a}19:{b}19", title, font=F(10, True, color), align=LEFT)
    for i, (lab, fml, bold) in enumerate(rows):
        r = 20 + i
        merge(ws, f"{a}{r}:{L(c2 - 2)}{r}", lab, font=F(9, bold, TEXT if not bold else INK), align=LEFT)
        merge(ws, f"{L(c2 - 1)}{r}:{b}{r}", fml, font=F(10, bold, INK), fmt=MONEY, align=RIGHT)
        if bold:
            for c in ws[f"{a}{r}:{b}{r}"][0]:
                c.border = Border(top=Side(style="thin", color=LINE_C))
    merge(ws, f"{a}25:{b}25", footnote, font=F(8, False, MUTED, True), align=LEFT_WRAP)
    ws.row_dimensions[25].height = 30


lens(2, 6, "FLUJO DE EFECTIVO (cuentas líquidas)", [
    ("Entradas de dinero", f'=SUMIFS(rFlujo,{FILT_D},rFlujo,">0")', False),
    ("Salidas de dinero", f'=-SUMIFS(rFlujo,{FILT_D},rFlujo,"<0")', False),
    ("Flujo neto", f"=SUMIFS(rFlujo,{FILT_D})", True),
    ("Pagos de tarjeta", f'=SUMIFS(rDOP,rTipoE,"Pago tarjeta",{FILT_D})', False),
], BLUE, "Efectivo, bancos y billeteras. Incluye pagos de tarjeta y depósitos a trading; excluye compras a crédito aún no pagadas.")
lens(7, 11, "GASTO REAL (consumo)", [
    ("Gastos fijos", f'=SUMIFS(rGasto,rClasE,"Fijo",{FILT_D},rCat,critCat,rCtaE,critCta)', False),
    ("Gastos variables", f'=SUMIFS(rGasto,rClasE,"Variable",{FILT_D},rCat,critCat,rCtaE,critCta)', False),
    ("Gasto real neto", f"={GASTO_P}", True),
    ("…de ello, con tarjeta de crédito", f'=SUMIFS(rGasto,rTCO,"Tarjeta de crédito",{FILT_D},rCat,critCat,rCtaE,critCta)', False),
], ORANGE, "Fecha de consumo, sin transferencias ni pagos de tarjeta. Reembolsos ya descontados.")
lens(12, 17, "TRADING (separado del consumo)", [
    ("Depósitos netos a trading", f'=SUMIFS(rDOP,rTipoE,"Transferencia",rTCD,"Cuenta de trading",{FILT_D})-SUMIFS(rDOP,rTipoE,"Transferencia",rTCO,"Cuenta de trading",{FILT_D})', False),
    ("Ganancia / pérdida realizada", f"=SUMIFS(rTrad,{FILT_D})", False),
    ("Costos operativos (evaluación, plataformas…)", f"=-SUMIFS(rGasto,rCat,catTrading,{FILT_D})", False),
    ("Resultado neto de trading", f"=SUMIFS(rTrad,{FILT_D})-SUMIFS(rGasto,rCat,catTrading,{FILT_D})", True),
], VIOLET, "Los depósitos son movimientos de capital, no pérdidas. Los costos operativos también cuentan en el gasto real.")
for r in range(20, 25):
    ws.row_dimensions[r].height = 17

# Barra de presupuesto
ws.row_dimensions[27].height = 26
merge(ws, "B27:Q27", "=budTexto", font=F(10, True, INK), fill_c=CARD, align=LEFT)
outline_box(ws, "B27:Q27")
ws.conditional_formatting.add("B27:Q27", FormulaRule(formula=["budNivel=3"], fill=fill(BAD_BG), font=Font(color=BAD, bold=True)))
ws.conditional_formatting.add("B27:Q27", FormulaRule(formula=["budNivel=2"], fill=fill(WARN_BG), font=Font(color=WARN, bold=True)))
ws.conditional_formatting.add("B27:Q27", FormulaRule(formula=["budNivel=1"], fill=fill(GOOD_BG), font=Font(color=GOOD, bold=True)))

# Distribución por categoría
section(ws, "B29:Q29", "¿En qué se va mi dinero?  ·  distribución por categoría")
ws.row_dimensions[29].height = 24
area(ws, "B30:Q45", fill_c=CARD)
outline_box(ws, "B30:Q45")
for j, h in enumerate(["Categoría", "", "", "Monto", "%"]):
    if h:
        put(ws, f"{L(2 + j)}31", h, F(8, True, MUTED), align=LEFT if j == 0 else RIGHT)
ws.merge_cells("B31:D31")
for j in range(12):
    r = 32 + j
    merge(ws, f"B{r}:D{r}", f"='Motor'!{S_CAT['col_name']}{7 + j}", font=F(9, False, INK), align=LEFT)
    put(ws, f"E{r}", f"='Motor'!{S_CAT['col_val']}{7 + j}", F(9, True, INK), fmt=NUM, align=RIGHT)
    put(ws, f"F{r}", f'=IF(OR(E{r}="",{GASTO_P.replace("rCat,critCat,", "")}=0),"",E{r}/{GASTO_P.replace("rCat,critCat,", "")})', F(9, False, MUTED), fmt=PCT, align=RIGHT)
ws.conditional_formatting.add("F32:F43", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="F4B795", showValue=True))
note(ws, "B44", "El gráfico ignora el filtro de categoría para mostrar el reparto completo.")
ch = BarChart()
ch.type = "bar"
ch.grouping = "clustered"
ch.gapWidth = 45
ch.add_data(Reference(MOT, range_string=S_CAT["val"]), titles_from_data=False)
ch.set_categories(Reference(MOT, range_string=S_CAT["name"]))
style_chart(ch, legend=False, gridlines=False)
ch.x_axis.scaling.orientation = "maxMin"
ch.y_axis.delete = True
color_series(ch.series[0], ORANGE)
data_labels(ch.series[0], '#,##0;-#,##0;""', "outEnd")
ch.width, ch.height = 18.5, 7.6
ws.add_chart(ch, "H30")

# Evolución mensual y semanal
section(ws, "B47:Q47", "Evolución  ·  ingresos vs gasto (12 meses) y gasto semanal (12 semanas)")
ws.row_dimensions[47].height = 24
area(ws, "B48:I63", fill_c=CARD)
outline_box(ws, "B48:I63")
area(ws, "J48:Q63", fill_c=CARD)
outline_box(ws, "J48:Q63")
ch = BarChart()
ch.type = "col"
ch.gapWidth = 60
ch.overlap = -10
ch.add_data(Reference(MOT, min_col=39, min_row=6, max_row=18), titles_from_data=True)  # Ingresos
ch.add_data(Reference(MOT, min_col=40, min_row=6, max_row=18), titles_from_data=True)  # Gasto real
ch.set_categories(Reference(MOT, min_col=38, min_row=7, max_row=18))
style_chart(ch)
color_series(ch.series[0], BLUE)
color_series(ch.series[1], ORANGE)
ch.y_axis.numFmt = "#,##0"
ch.width, ch.height = 15.8, 7.6
ws.add_chart(ch, "B48")
ch = LineChart()
ch.add_data(Reference(MOT, min_col=54, min_row=6, max_row=18), titles_from_data=True)
ch.set_categories(Reference(MOT, min_col=53, min_row=7, max_row=18))
style_chart(ch, legend=False)
color_series(ch.series[0], ORANGE, line=True)
ch.y_axis.numFmt = "#,##0"
ch.width, ch.height = 15.8, 7.6
ws.add_chart(ch, "J48")

# Métodos de pago y fijo vs variable
section(ws, "B65:Q65", "¿Cómo pago?  ·  gasto por cuenta o método  ·  fijos frente a variables")
ws.row_dimensions[65].height = 24
area(ws, "B66:I81", fill_c=CARD)
outline_box(ws, "B66:I81")
area(ws, "J66:Q81", fill_c=CARD)
outline_box(ws, "J66:Q81")
ch = BarChart()
ch.type = "bar"
ch.gapWidth = 50
ch.add_data(Reference(MOT, range_string=S_CTA["val"]), titles_from_data=False)
ch.set_categories(Reference(MOT, range_string=S_CTA["name"]))
style_chart(ch, legend=False, gridlines=False)
ch.x_axis.scaling.orientation = "maxMin"
ch.y_axis.delete = True
color_series(ch.series[0], BLUE)
data_labels(ch.series[0], '#,##0;-#,##0;""', "outEnd")
ch.width, ch.height = 15.8, 7.6
ws.add_chart(ch, "B66")
# Fijo vs variable (tabla + dona)
put(ws, "O68", "Fijos", F(9, True, INK), align=LEFT)
put(ws, "P68", "=J20", F(9, True, INK), fmt=NUM, align=RIGHT)
put(ws, "O69", "Variables", F(9, True, INK), align=LEFT)
put(ws, "P69", "=J21", F(9, True, INK), fmt=NUM, align=RIGHT)
put(ws, "O71", '=IF(P68+P69=0,"",FIXED(P69/(P68+P69)*100,0)&"% variable")', F(14, True, ORANGE), align=LEFT)
note(ws, "O72", "Lo variable es donde")
note(ws, "O73", "hay más margen de ajuste.")
ch = DoughnutChart(holeSize=62)
ch.add_data(Reference(ws, min_col=16, min_row=68, max_row=69), titles_from_data=False)
ch.set_categories(Reference(ws, min_col=15, min_row=68, max_row=69))
s = ch.series[0]
for i, col in enumerate([BLUE, ORANGE]):
    pt = DataPoint(idx=i)
    pt.graphicalProperties.solidFill = col
    pt.graphicalProperties.line.solidFill = "FFFFFF"
    s.dPt.append(pt)
s.dLbls = DataLabelList()
s.dLbls.showPercent = True
s.dLbls.showVal = False
s.dLbls.showCatName = False
s.dLbls.showSerName = False
s.dLbls.showLegendKey = False
s.dLbls.txPr = axis_text(900, "FFFFFF")
ch.legend.position = "b"
ch.legend.txPr = axis_text(900, TEXT)
try:
    ch.graphical_properties = GraphicalProperties(ln=LineProperties(noFill=True))
except Exception:
    pass
ch.width, ch.height = 9.5, 7.6
ws.add_chart(ch, "J66")

# Top 10 conceptos y comercios
section(ws, "B83:Q83", "Top 10  ·  conceptos (subcategorías) y comercios donde más gasto")
ws.row_dimensions[83].height = 24


def top_table(c1, c2, title, src, rows=10):
    a, b = L(c1), L(c2)
    area(ws, f"{a}84:{b}96", fill_c=CARD)
    outline_box(ws, f"{a}84:{b}96")
    merge(ws, f"{a}84:{L(c2 - 3)}84", title, font=F(8, True, MUTED), align=LEFT)
    put(ws, f"{L(c2 - 2)}84", "Monto", F(8, True, MUTED), align=RIGHT)
    put(ws, f"{L(c2 - 1)}84", "Veces", F(8, True, MUTED), align=RIGHT)
    put(ws, f"{b}84", "% gasto", F(8, True, MUTED), align=RIGHT)
    for j in range(rows):
        r = 85 + j
        put(ws, f"{a}{r}", j + 1, F(9, True, "94A3B8"), align=CENTER)
        merge(ws, f"{L(c1 + 1)}{r}:{L(c2 - 3)}{r}", f"='Motor'!{src['col_name']}{7 + j}", font=F(9, False, INK), align=LEFT)
        put(ws, f"{L(c2 - 2)}{r}", f"='Motor'!{src['col_val']}{7 + j}", F(9, True, INK), fmt=NUM, align=RIGHT)
        put(ws, f"{L(c2 - 1)}{r}", f"='Motor'!{src['col_cnt']}{7 + j}", F(9, False, MUTED), fmt=NUM, align=RIGHT)
        put(ws, f"{b}{r}", f'=IF(OR({L(c2 - 2)}{r}="",{GASTO_P}=0),"",{L(c2 - 2)}{r}/{GASTO_P})', F(9, False, MUTED), fmt=PCT, align=RIGHT)
    ws.conditional_formatting.add(f"{L(c2 - 2)}85:{L(c2 - 2)}94", DataBarRule(start_type="num", start_value=0, end_type="max", color="F4B795", showValue=True))


top_table(2, 9, "CONCEPTO", S_SUB)
top_table(10, 17, "COMERCIO / ESTABLECIMIENTO", S_COM)
note(ws, "B96", "Respeta los filtros de período, categoría y cuenta.")

# Comparación mensual
section(ws, "B98:Q98", "Comparación mensual  ·  12 meses hasta el mes de referencia")
ws.row_dimensions[98].height = 24
area(ws, "B99:Q112", fill_c=CARD)
outline_box(ws, "B99:Q112")
mhead = [("B", "C", "Mes"), ("D", "E", "Ingresos"), ("F", "G", "Gasto real"), ("H", "I", "Ahorro neto"), ("J", "J", "Tasa ahorro"),
         ("K", "L", "Fijos"), ("M", "N", "Variables"), ("O", "O", "Var. gasto"), ("P", "Q", "Comida preparada")]
for a, b, t in mhead:
    if a != b:
        ws.merge_cells(f"{a}99:{b}99")
    put(ws, f"{a}99", t, F(8, True, MUTED), align=RIGHT if t != "Mes" else LEFT)
srcmap = {"D": MC["Ingresos"], "F": MC["Gasto real"], "H": MC["Ahorro neto"], "K": MC["Fijos"], "M": MC["Variables"], "P": MC["Comida preparada"]}
for j in range(12):
    r = 100 + j
    mr = 18 - j  # más reciente arriba
    merge(ws, f"B{r}:C{r}", f"='Motor'!{MC['Etiqueta']}{mr}", font=F(9, True, INK), align=LEFT)
    for a, src in srcmap.items():
        b = L(column_index_from_string(a) + 1)
        merge(ws, f"{a}{r}:{b}{r}", f"='Motor'!{src}{mr}", font=F(9, a == "F", INK), fmt=NUM, align=RIGHT)
    put(ws, f"J{r}", f"='Motor'!{MC['Tasa ahorro']}{mr}", F(9, False, MUTED), fmt=PCT, align=RIGHT)
    if mr > 7:
        put(ws, f"O{r}", f"=IF(N('Motor'!{MC['Gasto real']}{mr - 1})=0,\"\",'Motor'!{MC['Gasto real']}{mr}/'Motor'!{MC['Gasto real']}{mr - 1}-1)",
            F(9, False, MUTED), fmt='+0%;[Red]-0%;0%', align=RIGHT)
ws.conditional_formatting.add("O100:O111", CellIsRule(operator="greaterThan", formula=["0.1"], font=Font(color=BAD, bold=True)))
ws.conditional_formatting.add("O100:O111", CellIsRule(operator="lessThan", formula=["-0.05"], font=Font(color=GOOD, bold=True)))
note(ws, "B112", "Var. gasto = cambio frente al mes anterior. Rojo: subió más de 10 %; verde: bajó más de 5 %.")
ws.freeze_panes = "A7"

# ============================================================================
# ALIMENTACIÓN
# ============================================================================
ws = ALI
LAST = 18
header(ws, "Análisis de alimentación", "Comer fuera vs. cocinar · con tus datos reales", LAST)
ws.column_dimensions["A"].width = 2
for c in range(2, LAST):
    ws.column_dimensions[L(c)].width = 11
ws.column_dimensions[L(LAST)].width = 2
area(ws, f"A3:{L(LAST)}130", fill_c=PAGE)
merge(ws, "B4:M4", '="Período: "&pTexto&"   ·   "&pDias&" días   (se cambia en el Dashboard)"', font=F(10, True, INK), align=LEFT)
link(ws, "N4", "Dashboard", "Cambiar período →", "B5")
ws["N4"].alignment = LEFT

FD = FILT_D
CPT = f"SUMIFS(rGasto,rGrupo,gCP,{FD})"
CPN = f'COUNTIFS(rGrupo,gCP,rTipoE,"Gasto",{FD})'
SIT = f"SUMIFS(rGasto,rGrupo,gSI,{FD})"
SIN = f'COUNTIFS(rGrupo,gSI,rTipoE,"Gasto",{FD})'
BST = f"SUMIFS(rGasto,rGrupo,gBS,{FD})"
ALT = f'SUMIFS(rGasto,rGrupo,"?*",{FD})'
kpi_rows(6)
kpi(ws, 6, 2, 5, "Comida preparada (fuera de casa)", f"={CPT}", MONEY,
    f'={CPN}&" compras  ·  ticket promedio RD$ "&IF({CPN}=0,0,FIXED({CPT}/{CPN},0))', ORANGE)
kpi(ws, 6, 6, 9, "Supermercado e ingredientes", f"={SIT}", MONEY,
    f'={SIN}&" compras  ·  ticket promedio RD$ "&IF({SIN}=0,0,FIXED({SIT}/{SIN},0))', BLUE)
kpi(ws, 6, 10, 13, "Bebidas, snacks y otros", f"={BST}", MONEY,
    f'="Total alimentación: RD$ "&FIXED({ALT},0)', YELLOW)
kpi(ws, 6, 14, 17, "% de alimentación comiendo fuera", f'=IF({ALT}=0,"–",{CPT}/{ALT})', '0%',
    f'="Alimentación = "&IF(SUMIFS(rGasto,{FD})=0,0,FIXED({ALT}/SUMIFS(rGasto,{FD})*100,0))&"% del gasto real"', ORANGE)
kpi_rows(11)


def avg_meal_formula(meal):
    # suma y cuenta de subcategorías marcadas con esa «Comida del día»
    return (f'=IF(SUMIF($E$38:$E$49,"{meal}",$G$38:$G$49)=0,"Sin datos",'
            f'SUMIF($E$38:$E$49,"{meal}",$H$38:$H$49)/SUMIF($E$38:$E$49,"{meal}",$G$38:$G$49))')


kpi(ws, 11, 2, 5, "Costo promedio por desayuno", avg_meal_formula("Desayuno"), MONEY,
    '=SUMIF($E$38:$E$49,"Desayuno",$G$38:$G$49)&" desayunos comprados"', ORANGE, 16)
kpi(ws, 11, 6, 9, "Costo promedio por almuerzo", avg_meal_formula("Almuerzo"), MONEY,
    '=SUMIF($E$38:$E$49,"Almuerzo",$G$38:$G$49)&" almuerzos comprados"', ORANGE, 16)
kpi(ws, 11, 10, 13, "Costo promedio por cena", avg_meal_formula("Cena"), MONEY,
    '=SUMIF($E$38:$E$49,"Cena",$G$38:$G$49)&" cenas compradas"', ORANGE, 16)
kpi(ws, 11, 14, 17, "Frecuencia de comida comprada", f"={CPN}/pDias*7", '0.0" / semana"',
    f'="Días con comida comprada: "&COUNTIF(\'Motor\'!{CALB["cp"]},">0")&" de "&pDias&"  ·  "&FIXED({CPN}/pDias*30.4,0)&" al mes"', ORANGE, 16)

# Comparativo comprar vs cocinar
section(ws, "B16:Q16", "Comer fuera vs. cocinar en casa  ·  comparación con tus datos (no se asume que cocinar sea más barato)")
ws.row_dimensions[16].height = 24
area(ws, "B17:Q27", fill_c=CARD)
outline_box(ws, "B17:Q27")
put(ws, "B18", "TUS DATOS (opcionales)", F(8, True, MUTED), align=LEFT)
rows_in = [
    (19, "Comidas que preparaste en casa en el período", None, "0"),
    (20, "% del súper e ingredientes que es comida (resto: limpieza, hogar…)", 1, "0%"),
    (21, "Costo estimado por comida en casa (solo si no tienes el conteo)", None, MONEY2),
]
for r, lab, val, fmt in rows_in:
    merge(ws, f"B{r}:G{r}", lab, font=F(9, False, TEXT), align=LEFT)
    ws.merge_cells(f"H{r}:I{r}")
    input_cell(ws, f"H{r}", val, fmt)
    for c in ws[f"H{r}:I{r}"][0]:
        c.border = INPUT_BOX
        c.fill = fill(INPUT_BG)
ws["H20"].comment = Comment("Supuesto inicial: 100 %. Ajústalo si en el súper también compras artículos que no son comida.", "Sistema")
add_name(wb, "aliCasaN", "'Alimentación'!$H$19")
add_name(wb, "aliPctSuper", "'Alimentación'!$H$20")
add_name(wb, "aliCasaEst", "'Alimentación'!$H$21")
put(ws, "K18", "RESULTADO", F(8, True, MUTED), align=LEFT)
cmp_rows = [
    (19, "Costo promedio por comida comprada", f'=IF({CPN}=0,"",{CPT}/{CPN})'),
    (20, "Costo por comida en casa", f'=IF(N(aliCasaN)>0,{SIT}*aliPctSuper/aliCasaN,IF(N(aliCasaEst)>0,aliCasaEst,""))'),
    (21, "Diferencia por comida (comprada − casa)", '=IF(OR(P19="",P20=""),"",P19-P20)'),
]
for r, lab, fml in cmp_rows:
    merge(ws, f"K{r}:O{r}", lab, font=F(9, r == 21, INK), align=LEFT)
    merge(ws, f"P{r}:Q{r}", fml, font=F(11, True, INK), fmt=MONEY, align=RIGHT)
add_name(wb, "aliCostoFuera", "'Alimentación'!$P$19")
add_name(wb, "aliCostoCasa", "'Alimentación'!$P$20")
merge(ws, "K23:Q26",
      '=IF(P19="","Aún no hay compras de comida preparada en el período.",IF(P20="","Para comparar, escribe cuántas comidas cocinaste en casa (o un costo estimado). Mientras tanto, cada comida comprada cuesta en promedio RD$ "&FIXED(P19,0)&".",IF(P21>0,"Según tus datos, una comida en casa cuesta RD$ "&FIXED(P21,0)&" menos ("&FIXED(P21/P19*100,0)&"% más barata) que una comprada.","Según tus datos, cocinar NO está saliendo más barato: cuesta RD$ "&FIXED(-P21,0)&" más por comida. Revisa qué compras en el súper.")))',
      font=F(10, True, INK), fill_c=SOFT, align=TOP_WRAP)
merge(ws, "B23:I26",
      "Cómo se calcula: costo por comida en casa = (súper + ingredientes del período × % que es comida) ÷ comidas preparadas en casa. "
      "Si cocinas para varias personas o te sobran ingredientes, ajusta el conteo o el %.",
      font=F(8, False, MUTED, True), align=TOP_WRAP)

# Simulador de ahorro
section(ws, "B29:Q29", "Simulador de ahorro  ·  ¿qué pasa si compro menos comida preparada?")
ws.row_dimensions[29].height = 24
area(ws, "B30:Q34", fill_c=CARD)
outline_box(ws, "B30:Q34")
sim_head = [("B", "C", "Reducción"), ("D", "E", "Comidas evitadas / semana"), ("F", "G", "Ahorro bruto / semana"),
            ("H", "I", "Costo de cocinarlas"), ("J", "K", "Ahorro neto / semana"), ("L", "M", "Ahorro neto / mes"),
            ("N", "O", "Ahorro neto / año"), ("P", "Q", "")]
for a, b, t in sim_head:
    ws.merge_cells(f"{a}30:{b}30")
    put(ws, f"{a}30", t, F(8, True, MUTED), align=CENTER)
ws.row_dimensions[30].height = 26
scen = [0.10, 0.25, 0.50]
ws.row_dimensions[31].height = 20
for i, p in enumerate(scen + [None]):
    r = 31 + i
    ws.merge_cells(f"B{r}:C{r}")
    if p is None:
        input_cell(ws, f"B{r}", 0.75, '0%" menos"')
    else:
        input_cell(ws, f"B{r}", p, '0%" menos"')
    for c in ws[f"B{r}:C{r}"][0]:
        c.border = INPUT_BOX
        c.fill = fill(INPUT_BG)
    freq = f"({CPN}/pDias*7)"
    merge(ws, f"D{r}:E{r}", f"={freq}*B{r}", font=F(10, False, INK), fmt="0.0", align=CENTER)
    merge(ws, f"F{r}:G{r}", f'=IF(aliCostoFuera="",0,D{r}*aliCostoFuera)', font=F(10, False, INK), fmt=MONEY, align=CENTER)
    merge(ws, f"H{r}:I{r}", f'=IF(aliCostoCasa="","Falta dato",D{r}*aliCostoCasa)', font=F(10, False, MUTED), fmt=MONEY, align=CENTER)
    merge(ws, f"J{r}:K{r}", f'=IF(ISNUMBER(H{r}),F{r}-H{r},"–")', font=F(10, True, INK), fmt=MONEY, align=CENTER)
    merge(ws, f"L{r}:M{r}", f'=IF(ISNUMBER(J{r}),J{r}*52/12,"–")', font=F(10, True, GOOD), fmt=MONEY, align=CENTER)
    merge(ws, f"N{r}:O{r}", f'=IF(ISNUMBER(J{r}),J{r}*52,"–")', font=F(11, True, GOOD), fmt=MONEY, align=CENTER)
    merge(ws, f"P{r}:Q{r}", f'=IF(ISNUMBER(J{r}),"","Bruto/año: RD$ "&FIXED(F{r}*52,0))', font=F(8, False, MUTED, True), align=CENTER)
note(ws, "B35", "Los % de reducción son editables (celdas amarillas). Sin costo en casa solo se muestra el ahorro bruto: el neto requiere ese dato.")

# Tabla por subcategoría de alimentación
section(ws, "B36:Q36", "Detalle por tipo de comida  ·  y gasto promedio por día de la semana")
ws.row_dimensions[36].height = 24
heads = [("B", "D", "Subcategoría"), ("E", "E", "Comida"), ("F", "F", "Grupo"), ("G", "G", "Nº compras"), ("H", "H", "Total"),
         ("I", "I", "Ticket prom."), ("J", "J", "% aliment."), ("K", "K", "Por semana")]
for a, b, t in heads:
    if a != b:
        ws.merge_cells(f"{a}37:{b}37")
    put(ws, f"{a}37", t, F(8, True, MUTED), CARD, align=LEFT if a == "B" else CENTER)
for j in range(12):
    r = 38 + j
    k = j + 1
    pos = f"MATCH({k},catFoodIdx,0)"
    sub = f'IFERROR(INDEX(catSub,{pos})&"","")'
    ws.merge_cells(f"B{r}:D{r}")
    put(ws, f"B{r}", f"={sub}", F(9, False, INK), CARD, align=LEFT)
    put(ws, f"E{r}", f'=IF(B{r}="","",IFERROR(INDEX(catComida,{pos})&"",""))', F(9, False, MUTED), CARD, align=CENTER)
    put(ws, f"F{r}", f'=IF(B{r}="","",IFERROR(INDEX(catGrupo,{pos})&"",""))', F(8, False, MUTED), CARD,
        align=Alignment(horizontal="left", vertical="center", shrink_to_fit=True))
    cat = f"INDEX(catCat,{pos})"
    put(ws, f"G{r}", f'=IF(B{r}="","",COUNTIFS(rCat,{cat},rSub,B{r},rTipoE,"Gasto",{FD}))', F(9, False, INK), CARD, NUM, CENTER)
    put(ws, f"H{r}", f'=IF(B{r}="","",SUMIFS(rGasto,rCat,{cat},rSub,B{r},{FD}))', F(9, True, INK), CARD, NUM, RIGHT)
    put(ws, f"I{r}", f'=IF(N(G{r})=0,"",H{r}/G{r})', F(9, False, INK), CARD, NUM, RIGHT)
    put(ws, f"J{r}", f'=IF(OR(B{r}="",{ALT}=0),"",H{r}/{ALT})', F(9, False, MUTED), CARD, PCT, RIGHT)
    put(ws, f"K{r}", f'=IF(B{r}="","",G{r}/pDias*7)', F(9, False, MUTED), CARD, "0.0", RIGHT)
ws.conditional_formatting.add("H38:H49", DataBarRule(start_type="num", start_value=0, end_type="max", color="F4B795", showValue=True))
outline_box(ws, "B37:K49")
area(ws, "L37:Q49", fill_c=CARD)
outline_box(ws, "L37:Q49")

# Gráfico por día de la semana (a la derecha de la tabla)
ch = BarChart()
ch.type = "col"
ch.gapWidth = 50
ch.add_data(Reference(MOT, min_col=66 + 6, min_row=6, max_row=13), titles_from_data=True)  # Promedio por día
ch.set_categories(Reference(MOT, min_col=67, min_row=7, max_row=13))
style_chart(ch, legend=False)
color_series(ch.series[0], ORANGE)
ch.y_axis.numFmt = "#,##0"
ch.width, ch.height = 11.6, 6.4
ws.add_chart(ch, "L37")

# Semanal y mensual
section(ws, "B51:Q51", "Comparación semanal y mensual  ·  comida preparada vs. súper e ingredientes")
ws.row_dimensions[51].height = 24
area(ws, "B52:I67", fill_c=CARD)
outline_box(ws, "B52:I67")
area(ws, "J52:Q67", fill_c=CARD)
outline_box(ws, "J52:Q67")
ch = BarChart()
ch.type = "col"
ch.gapWidth = 60
ch.add_data(Reference(MOT, min_col=55, min_row=6, max_row=18), titles_from_data=True)
ch.add_data(Reference(MOT, min_col=56, min_row=6, max_row=18), titles_from_data=True)
ch.set_categories(Reference(MOT, min_col=53, min_row=7, max_row=18))
style_chart(ch)
color_series(ch.series[0], ORANGE)
color_series(ch.series[1], BLUE)
ch.y_axis.numFmt = "#,##0"
ch.width, ch.height = 15.8, 7.6
ws.add_chart(ch, "B52")
ch = BarChart()
ch.type = "col"
ch.gapWidth = 60
ch.add_data(Reference(MOT, min_col=44, min_row=6, max_row=18), titles_from_data=True)
ch.add_data(Reference(MOT, min_col=45, min_row=6, max_row=18), titles_from_data=True)
ch.set_categories(Reference(MOT, min_col=38, min_row=7, max_row=18))
style_chart(ch)
color_series(ch.series[0], ORANGE)
color_series(ch.series[1], BLUE)
ch.y_axis.numFmt = "#,##0"
ch.width, ch.height = 15.8, 7.6
ws.add_chart(ch, "J52")

# Días con más gasto y día de la semana
section(ws, "B69:Q69", "¿Qué días gasto más en comida?")
ws.row_dimensions[69].height = 24
area(ws, "B70:H79", fill_c=CARD)
outline_box(ws, "B70:H79")
for a, t in (("B", "#"), ("C", "Fecha"), ("E", "Día"), ("F", "Alimentación"), ("H", "Comprada")):
    put(ws, f"{a}70", t, F(8, True, MUTED), align=CENTER if a in "BCE" else RIGHT)
ws.merge_cells("C70:D70")
ws.merge_cells("F70:G70")
for j in range(7):
    r = 71 + j
    put(ws, f"B{r}", j + 1, F(9, True, "94A3B8"), align=CENTER)
    merge(ws, f"C{r}:D{r}", f"='Motor'!{S_DAYS['date']}{7 + j}", font=F(9, True, INK), fmt=DATE, align=CENTER)
    put(ws, f"E{r}", f'=IF(C{r}="","",CHOOSE(WEEKDAY(C{r},2),"Lun","Mar","Mié","Jue","Vie","Sáb","Dom"))', F(9, False, MUTED), align=CENTER)
    merge(ws, f"F{r}:G{r}", f"='Motor'!{S_DAYS['val']}{7 + j}", font=F(9, True, INK), fmt=NUM, align=RIGHT)
    put(ws, f"H{r}", f"='Motor'!{S_DAYS['cp']}{7 + j}", F(9, False, MUTED), fmt=NUM, align=RIGHT)
note(ws, "B79", "Top 7 días del período.")
area(ws, "J70:Q79", fill_c=CARD)
outline_box(ws, "J70:Q79")
for a, b, t in (("J", "K", "Día"), ("L", "M", "Alimentación"), ("N", "O", "Promedio / día"), ("P", "Q", "Comidas compradas")):
    merge(ws, f"{a}70:{b}70", t, font=F(8, True, MUTED), align=RIGHT if a != "J" else LEFT)
for n in range(7):
    r = 71 + n
    mr = 7 + n
    merge(ws, f"J{r}:K{r}", f"='Motor'!{DC['Día']}{mr}", font=F(9, True, INK), align=LEFT)
    merge(ws, f"L{r}:M{r}", f"='Motor'!{DC['Alimentación']}{mr}", font=F(9, False, INK), fmt=NUM, align=RIGHT)
    merge(ws, f"N{r}:O{r}", f"='Motor'!{DC['Promedio por día']}{mr}", font=F(9, True, INK), fmt=NUM, align=RIGHT)
    merge(ws, f"P{r}:Q{r}", f"='Motor'!{DC['Nº comidas compradas']}{mr}", font=F(9, False, MUTED), fmt=NUM, align=RIGHT)
ws.conditional_formatting.add("N71:N77", DataBarRule(start_type="num", start_value=0, end_type="max", color="F4B795", showValue=True))
ws.freeze_panes = "A5"

# ============================================================================
# GASTOS HORMIGA
# ============================================================================
ws = HOR
LAST = 18
header(ws, "Gastos hormiga", "Pequeños, frecuentes y acumulables", LAST)
ws.column_dimensions["A"].width = 2
for c in range(2, LAST):
    ws.column_dimensions[L(c)].width = 11
ws.column_dimensions[L(LAST)].width = 2
area(ws, f"A3:{L(LAST)}90", fill_c=PAGE)
merge(ws, "B4:H4", "Gasto hormiga = gasto VARIABLE (sin costos de trading) menor o igual a:", font=F(10, True, INK), align=LEFT)
ws.merge_cells("I4:J4")
input_cell(ws, "I4", 500, MONEY)
for c in ws["I4:J4"][0]:
    c.border = INPUT_BOX
    c.fill = fill(INPUT_BG)
ws["I4"].comment = Comment("Parámetro de análisis (no es un límite de gasto). Ajústalo a tu realidad.", "Sistema")
add_name(wb, "Umbral", "'Gastos hormiga'!$I$4")
merge(ws, "K4:Q4", '="Período: "&pTexto&" · "&pDias&" días"', font=F(9, False, MUTED, True), align=RIGHT)
HT = f"SUMIFS(rGasto,rHorm,1,{FD})"
HN = f"COUNTIFS(rHorm,1,{FD})"
GT = f"SUMIFS(rGasto,{FD})"
kpi_rows(6)
kpi(ws, 6, 2, 5, "Nº de gastos hormiga", f"={HN}", NUM, f'="de "&COUNTIFS(rTipoE,"Gasto",{FD})&" gastos registrados"', ORANGE)
kpi(ws, 6, 6, 9, "Total acumulado", f"={HT}", MONEY, f'="Ticket promedio: RD$ "&IF({HN}=0,0,FIXED({HT}/{HN},0))', ORANGE)
kpi(ws, 6, 10, 13, "Peso sobre el gasto total", f'=IF({GT}=0,"–",{HT}/{GT})', "0.0%", f'="Gasto real del período: RD$ "&FIXED({GT},0)', ORANGE)
kpi(ws, 6, 14, 17, "Frecuencia", f"={HN}/pDias", '0.0" / día"', f'=FIXED({HN}/pDias*7,1)&" por semana  ·  "&FIXED({HN}/pDias*30.4,0)&" al mes"', ORANGE)
kpi_rows(11)
kpi(ws, 11, 2, 5, "Promedio diario", f"={HT}/pDias", MONEY, "Lo que se va cada día en pequeños consumos", BLUE, 16)
kpi(ws, 11, 6, 9, "Proyección mensual", f"={HT}/pDias*30.4", MONEY, "Si mantienes el ritmo actual", BLUE, 16)
kpi(ws, 11, 10, 13, "Proyección anual", f"={HT}/pDias*365", MONEY, "Si mantienes el ritmo actual", BLUE, 16)
kpi(ws, 11, 14, 17, "Si los reduces a la mitad", f"={HT}/pDias*365/2", MONEY, "Ahorro anual potencial estimado", GOOD, 16)

section(ws, "B16:Q16", "¿Dónde están los gastos hormiga?")
ws.row_dimensions[16].height = 24


def htable(c1, c2, title, src, r0=17):
    a, b = L(c1), L(c2)
    area(ws, f"{a}{r0}:{b}{r0 + 11}", fill_c=CARD)
    outline_box(ws, f"{a}{r0}:{b}{r0 + 11}")
    merge(ws, f"{a}{r0}:{L(c2 - 3)}{r0}", title, font=F(8, True, MUTED), align=LEFT)
    put(ws, f"{L(c2 - 2)}{r0}", "Total", F(8, True, MUTED), align=RIGHT)
    put(ws, f"{L(c2 - 1)}{r0}", "Veces", F(8, True, MUTED), align=RIGHT)
    put(ws, f"{b}{r0}", "% hormiga", F(8, True, MUTED), align=RIGHT)
    for j in range(10):
        r = r0 + 1 + j
        put(ws, f"{a}{r}", j + 1, F(9, True, "94A3B8"), align=CENTER)
        merge(ws, f"{L(c1 + 1)}{r}:{L(c2 - 3)}{r}", f"='Motor'!{src['col_name']}{7 + j}", font=F(9, False, INK), align=LEFT)
        put(ws, f"{L(c2 - 2)}{r}", f"='Motor'!{src['col_val']}{7 + j}", F(9, True, INK), fmt=NUM, align=RIGHT)
        put(ws, f"{L(c2 - 1)}{r}", f"='Motor'!{src['col_cnt']}{7 + j}", F(9, False, MUTED), fmt=NUM, align=RIGHT)
        put(ws, f"{b}{r}", f'=IF(OR({L(c2 - 2)}{r}="",{HT}=0),"",{L(c2 - 2)}{r}/{HT})', F(9, False, MUTED), fmt=PCT, align=RIGHT)
    ws.conditional_formatting.add(f"{L(c2 - 2)}{r0 + 1}:{L(c2 - 2)}{r0 + 10}", DataBarRule(start_type="num", start_value=0, end_type="max", color="F4B795", showValue=True))


htable(2, 9, "POR CATEGORÍA", S_HCAT)
htable(10, 17, "POR COMERCIO", S_HCOM)

section(ws, "B30:Q30", "Tendencia  ·  gasto hormiga por mes y por semana")
ws.row_dimensions[30].height = 24
area(ws, "B31:Q46", fill_c=CARD)
outline_box(ws, "B31:Q46")
ch = BarChart()
ch.type = "col"
ch.gapWidth = 60
ch.add_data(Reference(MOT, min_col=47, min_row=6, max_row=18), titles_from_data=True)
ch.set_categories(Reference(MOT, min_col=38, min_row=7, max_row=18))
style_chart(ch, legend=False)
color_series(ch.series[0], ORANGE)
ch.y_axis.numFmt = "#,##0"
ch.width, ch.height = 15.8, 7.6
ws.add_chart(ch, "B31")
ch = LineChart()
ch.add_data(Reference(MOT, min_col=57, min_row=6, max_row=18), titles_from_data=True)
ch.set_categories(Reference(MOT, min_col=53, min_row=7, max_row=18))
style_chart(ch, legend=False)
color_series(ch.series[0], ORANGE, line=True)
ch.y_axis.numFmt = "#,##0"
ch.width, ch.height = 15.8, 7.6
ws.add_chart(ch, "J31")
note(ws, "B47", "Mes (izquierda) y semana (derecha) usan el mes/semana de referencia del período elegido en el Dashboard.")

# ============================================================================
# RECURRENTES
# ============================================================================
ws = REC
header(ws, "Suscripciones y gastos recurrentes", "Registro informativo · no suma en los reportes", 22)
widths = {"A": 2, "B": 30, "C": 24, "D": 26, "E": 22, "F": 11, "G": 8, "H": 12, "I": 13, "J": 18, "K": 8, "L": 20,
          "M": 13, "N": 14, "O": 14, "P": 13, "Q": 10, "R": 14, "S": 26, "T": 6, "U": 6, "V": 2}
for col, w in widths.items():
    ws.column_dimensions[col].width = w
merge(ws, "B3:S4",
      "Este registro es una AGENDA: estima cuánto pagas en suscripciones y cuándo vence el próximo cobro. NO se suma en los reportes "
      "(evita duplicar). Cada cobro real se registra en «Registro» como Gasto. Usa el mismo nombre en «Comercio» para que la columna "
      "«Último cargo registrado» confirme que ya lo anotaste.",
      font=F(9, False, MUTED, True), align=TOP_WRAP)
ws.row_dimensions[3].height = 22
ws.row_dimensions[4].height = 22
merge(ws, "B5:L5", "DATOS (editable)", font=F(9, True, "1E3A8A"), fill_c="E8EEF6", align=LEFT)
merge(ws, "M5:S5", "CÁLCULO AUTOMÁTICO", font=F(9, True, MUTED), fill_c="EEF2F7", align=LEFT)
rh = {"B": "Servicio", "C": "Categoría", "D": "Subcategoría", "E": "Comercio (igual que en Registro)", "F": "Monto",
      "G": "Moneda", "H": "Frecuencia", "I": "Último cobro / inicio", "J": "Cuenta", "K": "¿Activa?", "L": "Notas",
      "M": "Monto DOP", "N": "Equivalente mensual", "O": "Equivalente anual", "P": "Próximo cobro", "Q": "Días",
      "R": "Último cargo registrado", "S": "Estado", "T": "m", "U": "d"}
for col, h in rh.items():
    put(ws, f"{col}6", h, F(9, True, "FFFFFF"), INK if col <= "L" and len(col) == 1 else "475569", align=CENTER)
ws.row_dimensions[6].height = 30
ws.column_dimensions["T"].hidden = True
ws.column_dimensions["U"].hidden = True
REC_DEMO = [
    ("DEMO · Herramienta de IA", "Tecnología y suscripciones", "Herramientas de inteligencia artificial", "Herramienta IA", 20, "USD", "Mensual", date(2026, 10, 3), "Tarjeta de crédito", "Sí"),
    ("DEMO · Streaming de video", "Tecnología y suscripciones", "Entretenimiento", "Streaming", 10.99, "USD", "Mensual", date(2026, 9, 18), "Tarjeta de crédito", "Sí"),
    ("DEMO · Internet del hogar", "Hogar y servicios", "Internet", "Proveedor de internet", 2050, "DOP", "Mensual", date(2026, 10, 5), "Banco principal", "Sí"),
    ("DEMO · Licencia de software", "Tecnología y suscripciones", "Software", "Software anual", 99, "USD", "Anual", date(2026, 9, 9), "Tarjeta de crédito", "Sí"),
    ("DEMO · Gimnasio", "Salud", "Gimnasio y bienestar", "Gimnasio", 2500, "DOP", "Mensual", date(2026, 10, 8), "Tarjeta de crédito", "Sí"),
]
for i in range(RECN):
    r = REC0 + i
    vals = REC_DEMO[i] if i < len(REC_DEMO) else (None,) * 10
    for col, v in zip("BCDEFGHIJK", vals):
        fmt = DATE if col == "I" else (MONEY2 if col == "F" else None)
        put(ws, f"{col}{r}", v, F(10, col == "B", INK), INPUT_BG, fmt, LEFT if col in "BCDEJ" else CENTER, BOX)
    put(ws, f"L{r}", "DEMOSTRACIÓN" if i < len(REC_DEMO) else None, F(9, False, MUTED), INPUT_BG, None, LEFT, BOX)
    calc = {
        "M": f'=IF(OR(B{r}="",F{r}=""),"",F{r}*IF(G{r}="USD",TasaUSD,1))',
        "N": f'=IF(OR(M{r}="",K{r}="No"),"",M{r}*IFERROR(INDEX(frecFactor,MATCH(H{r},frecNombre,0)),1))',
        "O": f'=IF(N{r}="","",N{r}*12)',
        "T": f'=IFERROR(INDEX(frecMeses,MATCH(H{r},frecNombre,0)),1)',
        "U": f'=IFERROR(INDEX(frecDias,MATCH(H{r},frecNombre,0)),0)',
        "P": (f'=IF(OR(B{r}="",NOT(ISNUMBER(I{r})),K{r}="No"),"",IF(I{r}>TODAY(),I{r},IF(T{r}>0,'
              f'IF(EDATE(I{r},T{r}*MAX(1,CEILING(DATEDIF(I{r},TODAY(),"m")/T{r},1)))<TODAY(),'
              f'EDATE(I{r},T{r}*(MAX(1,CEILING(DATEDIF(I{r},TODAY(),"m")/T{r},1))+1)),'
              f'EDATE(I{r},T{r}*MAX(1,CEILING(DATEDIF(I{r},TODAY(),"m")/T{r},1)))),'
              f'I{r}+MAX(1,U{r})*MAX(1,CEILING((TODAY()-I{r})/MAX(1,U{r}),1)))))'),
        "Q": f'=IF(P{r}="","",P{r}-TODAY())',
        "R": f'=IF(E{r}="","",IF(COUNTIFS(rCom,E{r},rTipoE,"Gasto")=0,"Sin registros",_xlfn.MAXIFS(rFecha,rCom,E{r},rTipoE,"Gasto")))',
        "S": (f'=IF(B{r}="","",IF(K{r}="No","Inactiva",IF(P{r}="","Falta fecha",IF(Q{r}<=0,"⚠ Vence hoy",'
              f'IF(Q{r}<=7,"⚠ Vence en "&Q{r}&" días","✓ Próximo en "&Q{r}&" días")))))'),
    }
    for col, fml in calc.items():
        fmt = {"M": MONEY2, "N": MONEY2, "O": MONEY2, "P": DATE, "Q": "0", "R": DATE}.get(col)
        put(ws, f"{col}{r}", fml, F(10, col in "NS", INK), SOFT, fmt, RIGHT if col in "MNO" else CENTER, BOX)
rt = RECL + 2
put(ws, f"B{rt}", "Totales (suscripciones activas)", F(10, True, INK))
put(ws, f"M{rt}", "Mensual:", F(9, True, MUTED), align=RIGHT)
put(ws, f"N{rt}", f"=SUM(N{REC0}:N{RECL})", F(11, True, INK), fmt=MONEY2, align=RIGHT)
put(ws, f"M{rt + 1}", "Anual:", F(9, True, MUTED), align=RIGHT)
put(ws, f"N{rt + 1}", f"=SUM(O{REC0}:O{RECL})", F(11, True, INK), fmt=MONEY2, align=RIGHT)
put(ws, f"P{rt}", "Vencen en ≤ 7 días:", F(9, True, MUTED), align=RIGHT)
put(ws, f"Q{rt}", f'=COUNTIFS(Q{REC0}:Q{RECL},">=0",Q{REC0}:Q{RECL},"<=7")', F(11, True, WARN), align=CENTER)
put(ws, f"P{rt + 1}", "Gasto fijo real este mes:", F(9, True, MUTED), align=RIGHT)
put(ws, f"Q{rt + 1}", '=SUMIFS(rGasto,rClasE,"Fijo",rFecha,">="&mesHoy,rFecha,"<="&EOMONTH(diaHoy,0))', F(11, True, INK), fmt=MONEY2, align=CENTER)
add_name(wb, "recMensual", f"'Recurrentes'!$N${rt}")
add_name(wb, "recProximos", f"'Recurrentes'!$Q${rt}")
ws.merge_cells(f"Q{rt + 1}:S{rt + 1}")
ws.conditional_formatting.add(f"S{REC0}:S{RECL}", FormulaRule(formula=[f'LEFT(S{REC0},1)="⚠"'], fill=fill(WARN_BG), font=Font(color=WARN, bold=True)))
ws.conditional_formatting.add(f"S{REC0}:S{RECL}", FormulaRule(formula=[f'LEFT(S{REC0},1)="✓"'], font=Font(color=GOOD)))
ws.conditional_formatting.add(f"B{REC0}:L{RECL}", FormulaRule(formula=[f'$L{REC0}="DEMOSTRACIÓN"'], fill=fill(DEMO_BG)))
for col, src in (("C", "=lstCat"), ("G", "=lstMonedas"), ("H", "=frecNombre"), ("J", "=lstCuentas"), ("K", "=lstSiNo")):
    dv_list(ws, src, f"{col}{REC0}:{col}{RECL}")
dv_list(ws, f"=OFFSET('Categorías'!$C${CAT0},MATCH($C{REC0},catCat,0)-1,0,COUNTIF(catCat,$C{REC0}),1)", f"D{REC0}:D{RECL}")
dv_list(ws, "=lstComercios", f"E{REC0}:E{RECL}", strict=False)
ws.freeze_panes = "C7"

# ============================================================================
# PRESUPUESTOS Y OBJETIVOS
# ============================================================================
ws = PRE
header(ws, "Presupuestos y objetivos", "Etapa 2 · se activa cuando escribes un límite", 15)
widths = {"A": 2, "B": 44, "C": 15, "D": 14, "E": 15, "F": 14, "G": 15, "H": 15, "I": 14, "J": 16, "K": 17, "L": 20, "M": 3, "N": 3, "O": 2}
for col, w in widths.items():
    ws.column_dimensions[col].width = w
merge(ws, "B3:L4",
      "Etapa 1: deja todo vacío y limítate a registrar. Etapa 2: cuando conozcas tus patrones (usa la columna «Prom. mensual (3 meses previos)» "
      "como referencia), escribe tus límites en las celdas amarillas. Un límite vacío = sin límite. El Dashboard mostrará el cumplimiento y las alertas.",
      font=F(9, False, MUTED, True), align=TOP_WRAP)
ws.row_dimensions[3].height = 22
ws.row_dimensions[4].height = 22
section(ws, "B6:L6", "Límites generales")
for j, h in enumerate(["Límite", "Monto límite", "Gasto actual", "% usado", "Disponible", "Estado"]):
    put(ws, f"{L(2 + j)}7", h, F(9, True, "FFFFFF"), INK if j < 2 else "475569", align=CENTER)
put(ws, "H7", "Alertar al llegar a", F(9, True, "FFFFFF"), INK, align=CENTER)
input_cell(ws, "H8", 0.8, "0%")
add_name(wb, "alertPct", "'Presupuestos'!$H$8")
gen = [
    (8, "Diario (diaHoy)", 'SUMIFS(rGasto,rFecha,diaHoy)'),
    (9, "Semanal (esta semana)", 'SUMIFS(rGasto,rFecha,">="&semHoy,rFecha,"<="&semHoy+6)'),
    (10, "Mensual (este mes)", 'SUMIFS(rGasto,rFecha,">="&mesHoy,rFecha,"<="&EOMONTH(diaHoy,0))'),
]
for r, lab, fml in gen:
    put(ws, f"B{r}", lab, F(10, True, INK), CARD, border=BOX, align=LEFT)
    input_cell(ws, f"C{r}", None, MONEY2)
    put(ws, f"D{r}", "=" + fml, F(10, False, INK), SOFT, MONEY2, RIGHT, BOX)
    put(ws, f"E{r}", f'=IF(N(C{r})=0,"",D{r}/C{r})', F(10, True, INK), SOFT, "0%", CENTER, BOX)
    put(ws, f"F{r}", f'=IF(N(C{r})=0,"",C{r}-D{r})', F(10, False, INK), SOFT, MONEY2, RIGHT, BOX)
    put(ws, f"G{r}", f'=IF(N(C{r})=0,"— Sin límite",IF(E{r}>1,"✕ Excedido",IF(E{r}>=alertPct,"⚠ Atención","✓ En rango")))', F(10, True, MUTED), SOFT, None, CENTER, BOX)
add_name(wb, "limDiario", "'Presupuestos'!$C$8")
add_name(wb, "limSemanal", "'Presupuestos'!$C$9")
add_name(wb, "limMensual", "'Presupuestos'!$C$10")
ws.conditional_formatting.add("E8:E10", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="9EC5F4", showValue=True))

section(ws, "B12:L12", "Presupuesto por categoría")
bh = ["Categoría", "Límite mensual", "Gasto este mes", "% del mes", "Estado mes", "Límite semanal", "Gasto esta semana",
      "Estado semana", "Prom. mensual (3 meses previos)", "Gasto mes anterior"]
for j, h in enumerate(bh):
    put(ws, f"{L(2 + j)}13", h, F(9, True, "FFFFFF"), INK if j in (0, 1, 5) else "475569", align=CENTER)
ws.row_dimensions[13].height = 30
cat_names = [c for c, _ in CATS if c != "Ingresos"]
B0, BN = 14, 30
BL = B0 + BN - 1
for i in range(BN):
    r = B0 + i
    put(ws, f"B{r}", cat_names[i] if i < len(cat_names) else None, F(10, True, INK), INPUT_BG, None, LEFT, BOX)
    input_cell(ws, f"C{r}", None, MONEY2)
    input_cell(ws, f"G{r}", None, MONEY2)
    mes = f'SUMIFS(rGasto,rCat,B{r},rFecha,">="&mesHoy,rFecha,"<="&EOMONTH(diaHoy,0))'
    sem = f'SUMIFS(rGasto,rCat,B{r},rFecha,">="&semHoy,rFecha,"<="&semHoy+6)'
    st = lambda lim, pct: f'IF(N({lim})=0,"—",IF({pct}>1,"✕ Excedido",IF({pct}>=alertPct,"⚠ Atención","✓ En rango")))'
    calc = {
        "D": f'=IF(B{r}="","",{mes})',
        "E": f'=IF(OR(B{r}="",N(C{r})=0),"",D{r}/C{r})',
        "F": f'=IF(B{r}="","",{st(f"C{r}", f"E{r}")})',
        "H": f'=IF(B{r}="","",{sem})',
        "I": f'=IF(B{r}="","",{st(f"G{r}", f"IF(N(G{r})=0,0,H{r}/G{r})")})',
        "J": f'=IF(B{r}="","",SUMIFS(rGasto,rCat,B{r},rFecha,">="&EDATE(mesHoy,-3),rFecha,"<"&mesHoy)/mesesDatos3)',
        "K": f'=IF(B{r}="","",SUMIFS(rGasto,rCat,B{r},rFecha,">="&EDATE(mesHoy,-1),rFecha,"<"&mesHoy))',
    }
    for col, fml in calc.items():
        fmt = {"D": MONEY2, "E": "0%", "H": MONEY2, "J": MONEY2, "K": MONEY2}.get(col)
        put(ws, f"{col}{r}", fml, F(10, col in "FI", INK if col not in "FI" else MUTED), SOFT, fmt, CENTER if col in "EFI" else RIGHT, BOX)
add_name(wb, "budLimMes", f"'Presupuestos'!$C${B0}:$C${BL}")
add_name(wb, "budLimSem", f"'Presupuestos'!$G${B0}:$G${BL}")
add_name(wb, "budEstMes", f"'Presupuestos'!$F${B0}:$F${BL}")
add_name(wb, "budEstSem", f"'Presupuestos'!$I${B0}:$I${BL}")
dv_list(ws, "=lstCat", f"B{B0}:B{BL}")
ws.conditional_formatting.add(f"E{B0}:E{BL}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="9EC5F4", showValue=True))
for rng in ("G8:G10", f"F{B0}:F{BL}", f"I{B0}:I{BL}"):
    first = rng.split(":")[0]
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT({first},1)="✕"'], fill=fill(BAD_BG), font=Font(color=BAD, bold=True)))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT({first},1)="⚠"'], fill=fill(WARN_BG), font=Font(color=WARN, bold=True)))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT({first},1)="✓"'], fill=fill(GOOD_BG), font=Font(color=GOOD, bold=True)))

# Resumen para el Dashboard
rs = BL + 2
put(ws, f"B{rs}", "Resumen (lo usa el Dashboard · no editar)", F(9, True, MUTED))
put(ws, f"B{rs + 1}", "Límites configurados", F(9, False, MUTED))
put(ws, f"C{rs + 1}", f"=COUNT(C8:C10)+COUNT(budLimMes)+COUNT(budLimSem)", F(9, True, INK))
put(ws, f"B{rs + 2}", "Alertas excedidas", F(9, False, MUTED))
put(ws, f"C{rs + 2}", '=COUNTIF(G8:G10,"✕*")+COUNTIF(budEstMes,"✕*")+COUNTIF(budEstSem,"✕*")', F(9, True, INK))
put(ws, f"B{rs + 3}", "Alertas en atención", F(9, False, MUTED))
put(ws, f"C{rs + 3}", '=COUNTIF(G8:G10,"⚠*")+COUNTIF(budEstMes,"⚠*")+COUNTIF(budEstSem,"⚠*")', F(9, True, INK))
put(ws, f"B{rs + 4}", "Nivel (0 sin límites · 1 ok · 2 atención · 3 excedido)", F(9, False, MUTED))
put(ws, f"C{rs + 4}", f"=IF(C{rs + 1}=0,0,IF(C{rs + 2}>0,3,IF(C{rs + 3}>0,2,1)))", F(9, True, INK))
put(ws, f"B{rs + 5}", "Texto", F(9, False, MUTED))
put(ws, f"C{rs + 5}",
    f'=IF(C{rs + 4}=0,"Etapa 1 · Sin límites configurados: sigue registrando. Cuando conozcas tus patrones, define límites en la hoja Presupuestos.",'
    f'"Presupuesto  ·  Mes: "&IF(N(limMensual)=0,"sin límite",FIXED(E10*100,0)&"% usado")&"   ·   Semana: "&IF(N(limSemanal)=0,"sin límite",FIXED(E9*100,0)&"% usado")'
    f'&"   ·   Excedidos: "&C{rs + 2}&"   ·   En atención: "&C{rs + 3})', F(9, False, INK))
add_name(wb, "budNivel", f"'Presupuestos'!$C${rs + 4}")
add_name(wb, "budTexto", f"'Presupuestos'!$C${rs + 5}")

# Objetivos
ro = rs + 8
section(ws, f"B{ro}:L{ro}", "Objetivos financieros")
oh = ["Objetivo", "Meta (DOP)", "Fecha objetivo", "Cuenta vinculada", "Aporte manual", "Avance actual", "% avance",
      "Falta", "Meses restantes", "Aporte mensual necesario", "Estado"]
for j, h in enumerate(oh):
    put(ws, f"{L(2 + j)}{ro + 1}", h, F(9, True, "FFFFFF"), INK if j in (0, 1, 2, 3, 4) else "475569", align=CENTER)
ws.row_dimensions[ro + 1].height = 30
goals = ["Fondo de emergencia", "Ahorro para la boda", "Inversión y crecimiento patrimonial", "Ahorro general"]
O0 = ro + 2
for i in range(10):
    r = O0 + i
    put(ws, f"B{r}", goals[i] if i < len(goals) else None, F(10, True, INK), INPUT_BG, None, LEFT, BOX)
    input_cell(ws, f"C{r}", None, MONEY2)
    input_cell(ws, f"D{r}", None, DATE)
    input_cell(ws, f"E{r}", None)
    input_cell(ws, f"F{r}", None, MONEY2)
    calc = {
        "G": f'=IF(B{r}="","",IF(E{r}<>"",IFERROR(INDEX(ctaSaldoDOP,MATCH(E{r},ctaNombre,0))+N(F{r}),N(F{r})),N(F{r})))',
        "H": f'=IF(OR(B{r}="",N(C{r})=0),"",MIN(1,N(G{r})/C{r}))',
        "I": f'=IF(OR(B{r}="",N(C{r})=0),"",MAX(0,C{r}-N(G{r})))',
        "J": f'=IF(OR(B{r}="",NOT(ISNUMBER(D{r}))),"",MAX(0,DATEDIF(MIN(TODAY(),D{r}),D{r},"m")))',
        "K": f'=IF(OR(I{r}="",J{r}=""),"",IF(J{r}=0,I{r},I{r}/J{r}))',
        "L": f'=IF(B{r}="","",IF(N(C{r})=0,"— Define una meta",IF(H{r}>=1,"✓ Cumplido",IF(AND(ISNUMBER(D{r}),D{r}<TODAY()),"✕ Fecha vencida","En progreso"))))',
    }
    for col, fml in calc.items():
        fmt = {"G": MONEY2, "H": "0%", "I": MONEY2, "J": "0", "K": MONEY2}.get(col)
        put(ws, f"{col}{r}", fml, F(10, col == "L", INK if col != "L" else MUTED), SOFT, fmt, CENTER if col in "HJL" else RIGHT, BOX)
OL = O0 + 9
dv_list(ws, "=lstCuentas", f"E{O0}:E{OL}")
ws.conditional_formatting.add(f"H{O0}:H{OL}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="8FD9B6", showValue=True))
ws.conditional_formatting.add(f"L{O0}:L{OL}", FormulaRule(formula=[f'LEFT(L{O0},1)="✓"'], fill=fill(GOOD_BG), font=Font(color=GOOD, bold=True)))
ws.conditional_formatting.add(f"L{O0}:L{OL}", FormulaRule(formula=[f'LEFT(L{O0},1)="✕"'], fill=fill(BAD_BG), font=Font(color=BAD, bold=True)))
add_name(wb, "objAvance", f"'Presupuestos'!$H${O0}:$H${OL}")

rf = OL + 2
section(ws, f"B{rf}:L{rf}", "Referencia para el fondo de emergencia")
put(ws, f"B{rf + 1}", "Meses de cobertura deseados", F(10, True, INK), align=LEFT)
input_cell(ws, f"C{rf + 1}", None, "0")
put(ws, f"B{rf + 2}", "Gasto mensual promedio (últimos 3 meses completos con datos)", F(10, False, TEXT), align=LEFT)
put(ws, f"C{rf + 2}", '=SUMIFS(rGasto,rFecha,">="&EDATE(mesHoy,-3),rFecha,"<"&mesHoy)/mesesDatos3', F(10, True, INK), fmt=MONEY2, align=RIGHT)
put(ws, f"B{rf + 3}", "Meta sugerida (meses × gasto promedio)", F(10, False, TEXT), align=LEFT)
put(ws, f"C{rf + 3}", f'=IF(N(C{rf + 1})=0,"Indica los meses",C{rf + 1}*C{rf + 2})', F(10, True, GOOD), fmt=MONEY2, align=RIGHT)
merge(ws, f"D{rf + 1}:L{rf + 3}", "Solo es una referencia calculada con tus gastos reales: copia el valor en «Meta» del Fondo de emergencia si te parece adecuado. "
      "Si vinculas una cuenta a un objetivo, el avance usa el saldo de esa cuenta (requiere saldo inicial) más el aporte manual.",
      font=F(9, False, MUTED, True), align=TOP_WRAP)
ws.freeze_panes = "C5"

# ============================================================================
# INICIO
# ============================================================================
ws = INI
LAST = 14
header(ws, "Finanzas Personales", "Excel 365 · sin macros · moneda base DOP", LAST)
ws.column_dimensions["A"].width = 2
for c in range(2, LAST):
    ws.column_dimensions[L(c)].width = 13
ws.column_dimensions[L(LAST)].width = 2
area(ws, f"A3:{L(LAST)}75", fill_c=PAGE)
merge(ws, "B4:M5", "Registra cada movimiento en segundos, entiende a dónde va tu dinero y, cuando tengas datos suficientes, define límites y metas. "
      "Todas las cifras se recalculan automáticamente al agregar filas en «Registro».", font=F(11, False, TEXT), align=TOP_WRAP)
ws.row_dimensions[4].height = 20
ws.row_dimensions[5].height = 20

tiles = [
    ("Registro", "Registrar movimientos", "Una fila por transacción. Fecha + monto + categoría + cuenta.", BLUE, "B6"),
    ("Dashboard", "Dashboard", "KPIs, filtros, gráficos, top 10 y comparación mensual.", BLUE, "A3"),
    ("Alimentación", "Alimentación", "Comer fuera vs. cocinar, costos por comida y simulador.", ORANGE, "A3"),
    ("Gastos hormiga", "Gastos hormiga", "Consumos pequeños y frecuentes y su peso real.", ORANGE, "A3"),
    ("Recurrentes", "Suscripciones", "Agenda de cobros recurrentes y próximos vencimientos.", ORANGE, "A6"),
    ("Presupuestos", "Presupuestos y metas", "Etapa 2: límites, alertas, fondo de emergencia, boda, inversión.", AQUA, "A3"),
    ("Cuentas", "Cuentas", "Saldos iniciales, saldo actual, conciliación y deuda de tarjetas.", AQUA, "A6"),
    ("Categorías", "Categorías", "Edita, agrega o elimina categorías y subcategorías.", MUTED, "A5"),
]
section(ws, "B7:M7", "Ir a…")
ws.row_dimensions[7].height = 24
for i, (sheet, title, desc, color, cell) in enumerate(tiles):
    row = 8 + (i // 4) * 4
    c1 = 2 + (i % 4) * 3
    a, b = L(c1), L(c1 + 2)
    area(ws, f"{a}{row}:{b}{row + 2}", fill_c=CARD)
    outline_box(ws, f"{a}{row}:{b}{row + 2}")
    for rr in range(row, row + 3):
        ws[f"{a}{rr}"].border = Border(left=Side(style="thick", color=color), top=ws[f"{a}{rr}"].border.top, bottom=ws[f"{a}{rr}"].border.bottom)
    ws.merge_cells(f"{a}{row}:{b}{row}")
    link(ws, f"{a}{row}", sheet, title + "  →", cell, font=F(11, True, INK))
    ws[f"{a}{row}"].alignment = Alignment(horizontal="left", vertical="bottom", indent=1)
    merge(ws, f"{a}{row + 1}:{b}{row + 2}", desc, font=F(9, False, MUTED), align=TOP_WRAP)
    ws.row_dimensions[row].height = 22
    ws.row_dimensions[row + 1].height = 18
    ws.row_dimensions[row + 2].height = 18

# Estado del sistema
section(ws, "B16:G16", "Estado del sistema")
section(ws, "I16:M16", "Primeros pasos")
ws.row_dimensions[16].height = 24
status = [
    ("Movimientos registrados", '=COUNTIF(rTipoE,"?*")', NUM),
    ("Filas de DEMOSTRACIÓN", '=COUNTIF(rObs,"DEMOSTRACIÓN")', NUM),
    ("Filas con advertencias", '=COUNTIF(rEstado,"⚠*")', NUM),
    ("Último movimiento", '=IF(COUNT(rFecha)=0,"—",MAX(rFecha))', DATE),
    ("Capacidad usada del Registro", f'=COUNTA(rFecha)/{NROWS}', "0.0%"),
    ("Cuentas con saldo inicial", "='Cuentas'!L" + str(r_tot + 3), None),
    ("Tasa USD→DOP en uso", "=TasaUSD", "0.00"),
    ("Suscripciones que vencen en ≤ 7 días", "=recProximos", NUM),
]
for i, (lab, fml, fmt) in enumerate(status):
    r = 17 + i
    merge(ws, f"B{r}:E{r}", lab, font=F(10, False, TEXT), fill_c=CARD, align=LEFT, border=BOTTOM)
    merge(ws, f"F{r}:G{r}", fml, font=F(10, True, INK), fill_c=CARD, fmt=fmt, align=RIGHT)
    ws.row_dimensions[r].height = 20
ws.conditional_formatting.add("F19", CellIsRule(operator="greaterThan", formula=["0"], font=Font(color=BAD, bold=True), fill=fill(BAD_BG)))
ws.conditional_formatting.add("F18", CellIsRule(operator="greaterThan", formula=["0"], font=Font(color=VIOLET, bold=True), fill=fill(DEMO_BG)))
steps = [
    "1. Explora el Dashboard con los datos DEMO (filas lilas del Registro).",
    "2. Borra los DEMO (ver «Eliminar datos de demostración» abajo).",
    "3. En Cuentas, renombra tus cuentas/tarjetas y, si quieres saldos reales, escribe saldo inicial y su fecha.",
    "4. En Config, actualiza la tasa USD→DOP.",
    "5. Ajusta Categorías a tu gusto (opcional).",
    "6. Registra cada gasto el mismo día: ~10 segundos por fila.",
    "7. Tras 2–3 meses, revisa Alimentación y Hormiga y define límites en Presupuestos.",
]
for i, t in enumerate(steps):
    merge(ws, f"I{17 + i}:M{17 + i}", t, font=F(9, False, TEXT), align=LEFT_WRAP)

section(ws, "B26:M26", "Cómo registrar (reglas que evitan duplicar)")
ws.row_dimensions[26].height = 24
rules = [
    ("Gasto cotidiano", "Fecha · (Tipo vacío = Gasto) · Categoría · Subcategoría · Monto · Cuenta. Comercio y descripción son opcionales."),
    ("Compra con tarjeta de crédito", "Regístrala como Gasto en la fecha de consumo, con Cuenta = la tarjeta. Así el gasto cuenta una sola vez y la deuda de la tarjeta crece."),
    ("Pago de la tarjeta", "Tipo «Pago tarjeta», Cuenta = banco que paga, Cuenta destino = la tarjeta. NO es gasto: solo reduce la deuda y sale dinero del banco."),
    ("Transferencias propias", "Tipo «Transferencia» (banco→efectivo, banco→billetera…). No suma ni a ingresos ni a gastos; solo mueve saldos."),
    ("Trading", "Depósito o retiro = Transferencia hacia/desde la cuenta de trading. Ganancia o pérdida realizada = «Resultado trading» (+/−). "
                "Evaluaciones, plataformas y comisiones = Gasto en la categoría de trading."),
    ("Reembolsos", "Si te devuelven dinero, usa «Reembolso» con la misma categoría del gasto original: el gasto neto se reduce."),
    ("Dólares", "Moneda USD. Si conoces la tasa real aplicada, escríbela en «Tasa USD»; si no, se usa la de Config. Todos los reportes salen en DOP."),
    ("Ajustes", "Si el saldo real no coincide con el calculado (Cuentas › Conciliación), registra un «Ajuste» por la diferencia."),
]
for i, (t, d) in enumerate(rules):
    r = 27 + i
    merge(ws, f"B{r}:D{r}", t, font=F(10, True, INK), fill_c=CARD, align=LEFT_WRAP)
    merge(ws, f"E{r}:M{r}", d, font=F(9, False, TEXT), fill_c=CARD, align=LEFT_WRAP)
    ws.row_dimensions[r].height = 32
    for c in ws[f"B{r}:M{r}"][0]:
        c.border = BOTTOM

section(ws, "B36:M36", "Eliminar datos de demostración")
ws.row_dimensions[36].height = 24
demo_txt = [
    f"Las filas DEMO están juntas al inicio del Registro (filas {R0} a {R0 + len(DEMO) - 1}), marcadas «DEMOSTRACIÓN» en Observaciones y con fondo lila.",
    f"Opción rápida: en Registro selecciona B{R0}:N{R0 + len(DEMO) - 1} y pulsa Suprimir. Las fórmulas automáticas (columnas O en adelante) no se tocan.",
    "En Recurrentes, borra igual las filas marcadas «DEMOSTRACIÓN» (columnas B:L).",
    "Las cuentas y categorías de ejemplo NO son datos personales: renómbralas o bórralas. No hay saldos ni límites precargados.",
    "Cuando el estado de arriba muestre 0 filas DEMO, todo el sistema queda listo con tus datos.",
]
for i, t in enumerate(demo_txt):
    merge(ws, f"B{37 + i}:M{37 + i}", "•  " + t, font=F(9, False, TEXT), align=LEFT_WRAP)
    ws.row_dimensions[37 + i].height = 20

section(ws, "B43:M43", "Buenas prácticas")
ws.row_dimensions[43].height = 24
tips_txt = [
    "No insertes columnas dentro de la tabla del Registro ni escribas sobre las columnas grises: son el motor del sistema.",
    f"El Registro tiene capacidad para {NROWS:,} filas. Si te acercas al límite, selecciona la última fila completa y arrastra el controlador de la tabla hacia abajo para ampliarla.",
    "Puedes ordenar y filtrar la tabla (flechas del encabezado) sin romper nada. Filtra la columna Estado por «⚠» para corregir errores.",
    "Celdas amarillas = editables. Celdas grises = cálculo automático. Filas lilas = demostración.",
    "En el celular: abre el archivo con la app de Excel (OneDrive) y registra desde la hoja Registro.",
]
for i, t in enumerate(tips_txt):
    merge(ws, f"B{44 + i}:M{44 + i}", "•  " + t, font=F(9, False, TEXT), align=LEFT_WRAP)
    ws.row_dimensions[44 + i].height = 20

# ----------------------------------------------------------------------------
# Ajustes finales
# ----------------------------------------------------------------------------
for w in wb.worksheets:
    nav(w)
    w.sheet_view.zoomScale = 90 if w.title in ("Registro", "Motor", "Config") else 100
MOT.sheet_view.zoomScale = 80
wb.active = 0
wb.calculation.fullCalcOnLoad = True
import os
if os.environ.get("PREVIEW") == "reg":
    for w in wb.worksheets:
        if w.title != "Registro":
            w.sheet_state = "hidden"
    REG.print_area = "A1:O45"
    REG.page_setup.orientation = "landscape"
    REG.page_setup.fitToWidth = 1
    REG.page_setup.fitToHeight = 0
    REG.sheet_properties.pageSetUpPr.fitToPage = True
elif os.environ.get("PREVIEW"):
    for nm in ("Registro", "Motor", "Config", "Categorías"):
        wb[nm].sheet_state = "hidden"
for nm in ("Inicio", "Dashboard", "Alimentación", "Gastos hormiga", "Presupuestos"):
    w = wb[nm]
    w.page_setup.orientation = "landscape" if nm != "Inicio" else "portrait"
    w.page_setup.fitToWidth = 1
    w.page_setup.fitToHeight = 0
    w.sheet_properties.pageSetUpPr.fitToPage = True
    w.print_options.horizontalCentered = True
    w.page_margins.left = w.page_margins.right = 0.3
for nm, rng in (("Cuentas", f"A1:U{CTAL + 6}"), ("Recurrentes", f"A1:S{RECL + 3}")):
    w = wb[nm]
    w.print_area = rng
    w.page_setup.orientation = "landscape"
    w.page_setup.fitToWidth = 1
    w.page_setup.fitToHeight = 0
    w.sheet_properties.pageSetUpPr.fitToPage = True
wb.save(OUT)
print("OK ->", OUT, "| filas DEMO:", len(DEMO))
